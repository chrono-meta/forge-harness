#!/usr/bin/env python3
"""test_lane_speech_coach.py — L16 speech-coach known-pair · 퇴화 · 되돌림.

🟥 픽스처는 전부 **실행 중에 합성**한다(톱니형 배음 + 결정적 LCG 잡음). 녹화 원본·사람 목소리·
   발표 내용은 넣지 않는다 — 공개 자산에 현장 녹음을 남기지 않는 규율이다. 받아쓰기·디코더는
   «가짜 도구」(이 파일이 임시 폴더에 쓰는 작은 python 스크립트)로 대신해서, whisper 나 ffmpeg 가
   없는 기계(CI)에서도 파이프라인 전 구간이 돈다. 진짜 ffmpeg 가 있으면 한 팔을 더 돈다(없으면
   SKIP — **SKIP 은 PASS 가 아니다**, 수를 따로 찍는다).

    python3 test_lane_speech_coach.py      # 0 통과 / 1 실패
"""
import array
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import wave

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import lane_speech_coach as L   # noqa: E402

SR = L.SR
FAILS, PASSES, SKIPS = [], [], []


def says_pass(text):
    """렌더 글자에 «통과」/PASS 가 있나 — 부정문 «통과 아님」은 뺀다."""
    t = text.replace('통과 아님', '')
    return '통과' in t or re.search(r'\bPASS\b', t) is not None


def chk(name, cond, extra=''):
    print('  %s %s %s' % ('✅' if cond else '❌', name, extra))
    (PASSES if cond else FAILS).append(name)


def skip(name, why):
    print('  ⏭  %s — SKIP (통과 아님): %s' % (name, why))
    SKIPS.append(name)


# ── 합성기 ──────────────────────────────────────────────────────────────────────────────────
class Synth:
    def __init__(self, seed=7):
        self.x = []
        self.s = seed

    def _noise(self, n, amp):
        out, s = [], self.s
        for _ in range(n):
            s = (1103515245 * s + 12345) & 0x7fffffff
            out.append(int((s / 0x7fffffff * 2 - 1) * amp))
        self.s = s
        return out

    def sil(self, dur, namp=20):
        self.x += self._noise(int(dur * SR), namp)
        return self

    def tone(self, dur, st_fn=lambda u: 0.0, f0=150.0, amp=8000, namp=20, harmonics=5):
        """st_fn(u): 0..1 진행률 → 반음 오프셋. 배음 1..harmonics(1/k) · 10ms 경사."""
        n = int(dur * SR)
        nz = self._noise(n, namp)
        ramp = int(0.01 * SR)
        ph = 0.0
        norm = sum(1.0 / k for k in range(1, harmonics + 1))
        for i in range(n):
            f = f0 * 2 ** (st_fn(i / n) / 12)
            ph += 2 * math.pi * f / SR
            v = sum(math.sin(k * ph) / k for k in range(1, harmonics + 1)) / norm
            env = min(1.0, i / ramp, (n - 1 - i) / ramp)
            self.x.append(int(amp * v * env) + nz[i])
        return self

    def arr(self):
        return array.array('h', [max(-32768, min(32767, v)) for v in self.x])

    def wav(self, path):
        L.write_wav(path, self.arr())
        return path


def tail_fn(kind, body=1.2, tail=0.3, depth=3.0):
    """본문 평탄 → 끝 tail 초 동안 fall/rise/flat/pushpull."""
    tot = body + tail

    def f(u):
        t = u * tot
        if t < body:
            return 0.0
        v = (t - body) / tail
        return {'fall': -depth * v, 'rise': depth * v, 'flat': 0.0,
                'pushpull': depth * (1 - abs(2 * v - 1))}[kind]
    return f


# ── 가짜 도구 ───────────────────────────────────────────────────────────────────────────────
FAKE_FFMPEG = r'''
import sys, wave
a = sys.argv
src = a[a.index('-i') + 1]
try:
    w = wave.open(src, 'rb')
except Exception:
    sys.stderr.write('fake-ffmpeg: not a wav\n'); sys.exit(1)
if w.getnchannels() != 1 or w.getframerate() != 16000 or w.getsampwidth() != 2:
    sys.stderr.write('fake-ffmpeg: need 16k mono s16\n'); sys.exit(1)
import os
n = w.getnframes()
keep = float(os.environ.get('FAKE_FFMPEG_KEEP', '1'))   # 디코드된 몫(잘림 재현)
err = os.environ.get('FAKE_FFMPEG_ERR')                 # rc 0 인데 stderr 에 오류(손상 패킷 재현)
if err:
    sys.stderr.write(err + '\n')
sys.stdout.buffer.write(w.readframes(int(n * keep)))
'''

FAKE_FFPROBE = r'''
import sys, wave
src = [x for x in sys.argv[1:] if not x.startswith('-') and '=' not in x and x not in ('a:0', 'error')][-1]
w = wave.open(src, 'rb')
d = w.getnframes() / w.getframerate()
print('duration=%.6f' % d)
print('duration=%.6f' % d)
'''

FAKE_WHISPER = r'''
import sys, os, json, re, wave
texts = json.loads(os.environ.get('FAKE_WHISPER_TEXTS', '[]'))
skip = {int(v) for v in os.environ.get('FAKE_WHISPER_SKIP', '').split(',') if v.strip()}
for f in [x for x in sys.argv[1:] if x.endswith('.wav')]:
    j = int(re.search(r'c(\d+)\.wav$', f).group(1))
    if j in skip:
        continue          # 이 덩어리는 JSON 을 안 쓴다(받아쓰기 누락 재현)
    raw = json.loads(os.environ.get('FAKE_WHISPER_RAW', '{}'))
    if str(j) in raw:     # 모양이 이상한 JSON 을 그대로 쓴다([] · {"transcription": null} …)
        open(f + '.json', 'w').write(raw[str(j)]); continue
    t = texts[j] if j < len(texts) else ''
    w = wave.open(f, 'rb'); dur = w.getnframes() / w.getframerate(); w.close()
    ws = t.split(); toks = [{"text": "[_BEG_]", "offsets": {"from": 0, "to": 0}, "p": 1.0}]
    a0, a1 = 0.12, max(0.13, dur - 0.12)
    for k, x in enumerate(ws):
        s = a0 + (a1 - a0) * k / max(1, len(ws)); e = a0 + (a1 - a0) * (k + 1) / max(1, len(ws))
        toks.append({"text": " " + x, "offsets": {"from": int(s * 1000), "to": int(e * 1000)}, "p": 0.9})
    json.dump({"transcription": [{"text": t, "tokens": toks}]}, open(f + '.json', 'w'), ensure_ascii=False)
'''


def install_tool(d, name, body):
    p = os.path.join(d, name)
    with open(p, 'w') as fh:
        fh.write('#!' + sys.executable + '\n' + body)
    os.chmod(p, 0o755)
    return p


# 리허설 픽스처: 덩어리 8개(4번은 0.2초 잡음 — 가짜 받아쓰기가 «감사합니다」를 지어낸다)
GREET_TEXTS = ['안녕하세요.', '오늘은 놓친 것을 이야기합니다.', '그런데', '답이 달라지는 곳이 있습니다.',
               '감사합니다.', '고치는 법을 봅니다.', '다음으로 넘어갑니다.', '감사합니다.']
GREET_SCRIPT = '안녕하세요.\n오늘은 놓치는 것을 이야기합니다.\n[진행] 다음 장\n그런데 답이 달라지는 곳이 있습니다.\n' \
               '고치는 법을 봅니다.\n다음으로 넘어갑니다.\n감사합니다.\n'


def greet_fixture(path, blip='noise'):
    """blip: 4번 덩어리(0.2초) — 'noise' = 무성 잡음 폭발(숨·마찰) · 'voiced' = 짧은 유성 발화(«네»)."""
    s = Synth(11).sil(0.5)
    s.tone(1.0, tail_fn('fall', 0.7, 0.3), amp=11000).sil(0.7)
    s.tone(1.6, tail_fn('fall', 1.3, 0.3)).sil(0.7)
    s.tone(0.5, amp=8000).sil(0.7)                        # 홀로 선 «그런데»
    s.tone(1.6, tail_fn('fall', 1.3, 0.3)).sil(0.7)
    if blip == 'voiced':
        s.tone(0.2, amp=3000).sil(0.7)                    # 짧은 «유성» 덩어리(실제 발화일 수 있다)
    else:
        s.sil(0.2, namp=3000).sil(0.7)                    # 짧은 «무성» 잡음 덩어리
    s.tone(1.5, tail_fn('pushpull', 1.2, 0.3)).sil(0.7)
    s.tone(1.5, tail_fn('rise', 1.2, 0.3)).sil(0.7)
    s.tone(0.9, tail_fn('fall', 0.6, 0.3), amp=1800).sil(0.6)   # 작은 끝인사
    return s.wav(path)


def main():
    tmp = tempfile.mkdtemp(prefix='speechcoach_pair_')
    try:
        run(tmp)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print('PASS %d · FAIL %d · SKIP %d (SKIP != PASS)' % (len(PASSES), len(FAILS), len(SKIPS)))
    if FAILS:
        print('FAIL · ' + ', '.join(FAILS))
    return 1 if FAILS else 0


def run(tmp):
    print('L16 speech-coach known-pair')
    # K1 쉼: 0.5 · 1.0 초 쉼은 경계, 0.10초는 흡수
    s = Synth(1).sil(0.5).tone(1.0).sil(0.5).tone(1.0).sil(1.0).tone(1.0).sil(0.10).tone(1.0).sil(0.5)
    db = L.frame_db(s.arr())
    ch, _ = L.segment(db)
    gaps = [round(ch[i + 1][0] - ch[i][1], 2) for i in range(len(ch) - 1)]
    chk('K1 쉼 경계: 덩어리 3개(0.10초 쉼은 흡수)', len(ch) == 3, 'chunks=%d gaps=%s' % (len(ch), gaps))
    chk('K1-b 쉼 길이 0.5 · 1.0초를 ±0.05 안으로', len(gaps) == 2 and abs(gaps[0] - 0.5) <= 0.05
        and abs(gaps[1] - 1.0) <= 0.05, str(gaps))

    # K2 문장 끝 피치 — 같은 본문, 끝 0.3초만 다르다(한 변수)
    res = {}
    for kind in ('fall', 'rise', 'flat', 'pushpull'):
        s = Synth(2).sil(0.4).tone(1.5, tail_fn(kind)).sil(0.6)
        x = s.arr()
        db = L.frame_db(x)
        ch, thr = L.segment(db)
        res[kind] = L.end_metrics(x, db, thr, ch[0][0], ch[0][1]) if len(ch) == 1 else {}
    ts = {k: v.get('tail_st') for k, v in res.items()}
    chk('K2 컨트롤 생존: 네 팔 모두 끝 피치를 잼', all(v is not None for v in ts.values()), str(ts))
    if all(v is not None for v in ts.values()):
        chk('K2-a 하강 → 음수(<-0.8)', ts['fall'] < -0.8, '%+.1f' % ts['fall'])
        chk('K2-b 상승 → 양수(>+0.8)', ts['rise'] > 0.8, '%+.1f' % ts['rise'])
        chk('K2-c 평탄 → |x|<0.5', abs(ts['flat']) < 0.5, '%+.1f' % ts['flat'])
        pp = {k: L.is_push_pull(v) for k, v in res.items()}
        chk('K2-d 밀당은 «올렸다 내림」에서만', pp == {'fall': False, 'rise': False, 'flat': False,
                                               'pushpull': True}, str(pp))

    # K3 끝인사 크기 — 끝 덩어리만 −14dB
    def greet_run(last_amp):
        s = Synth(3).sil(0.4)
        for _ in range(11):
            s.tone(0.8, harmonics=1).sil(0.5)
        s.tone(0.8, amp=last_amp, harmonics=1).sil(0.5)
        db = L.frame_db(s.arr())
        ch, _ = L.segment(db)
        return len(ch), L.row_greet(db, ch, None)
    n_q, g_q = greet_run(1600)
    n_c, g_c = greet_run(8000)
    chk('K3 컨트롤 생존: 양팔 덩어리 12개', n_q == 12 and n_c == 12, '%d/%d' % (n_q, n_c))
    chk('K3-a 작은 끝인사 → 하위 10% 아래', g_q['last']['below_p10'], '%+.1f dB' % g_q['last']['vs_median'])
    chk('K3-b 같은 크기 끝인사 → 하위 10% 아래 아님', not g_c['last']['below_p10'],
        '%+.1f dB' % g_c['last']['vs_median'])

    # K4 접속사 늘임 — 홀로 선 «그런데」 길이만 다르다
    def conj_run(conj_dur):
        ch = [(0.0, 1.0), (1.5, 2.5), (3.0, 3.0 + conj_dur), (4.0, 5.0), (5.5, 6.5)]
        ws = []
        for j, (a, b) in enumerate(ch):
            ws.append({'w': '그런데' if j == 2 else ('가나다라마바,' if j == 0 else '가나다라마바'),
                       't0': a, 't1': b, 'chunk': j, 'p': 1.0})
        return L.row_conj(ws, ch)
    r_n, r_d = conj_run(0.5), conj_run(0.8)
    sn, sd = r_n['conj'][0]['stretch'], r_d['conj'][0]['stretch']
    chk('K4-a 보통 «그런데」 → 늘임 ≈1.0, 끌기 문턱 아래', sn is not None and abs(sn - 1.0) <= 0.05
        and sn < L.C('drag_stretch'), str(sn))
    chk('K4-b 끈 «그런데」 → 끌기 문턱 이상', sd is not None and sd >= L.C('drag_stretch'), str(sd))
    chk('K4-c 홀로 선 접속사 길이는 «음향 경계」에서 잰다(토큰 시각 아님)',
        r_n['conj'][0]['dur_src'] == '음향 경계' and r_n['conj'][0]['pause_before'] == 0.5, str(r_n['conj'][0]))
    chk('K4-d 쉼표 뒤 쉼 실현', r_n['commas'] == [0.5], str(r_n['commas']))
    ch_j = [(0.0, 1.0), (1.5, 1.8), (2.3, 3.3)]
    ws_j = [{'w': w, 't0': a, 't1': b, 'chunk': j, 'p': 1.0}
            for j, ((a, b), w) in enumerate(zip(ch_j, ['가나다라마바', '즉', '가나다라마바']))]
    rj = L.row_conj(ws_j, ch_j)['conj'][0]
    chk('K4-e 1음절 «즉」은 늘임을 안 낸다(None — 0 이 아니다)', rj['conj'] == '즉' and rj['stretch'] is None, str(rj))
    chk('K4-f 옥타브 의심 끝(+13반음 봉우리)은 밀당으로 안 센다',
        not L.is_push_pull({'tail_st': -7.0, 'tail_peak_st': 13.2, 'tail_last_st': -7.2})
        and L.is_push_pull({'tail_st': 1.0, 'tail_peak_st': 2.5, 'tail_last_st': 0.5}))
    chk('K4-g 정상적인 깊은 평서 하강(−7.5반음)은 옥타브 의심이 아니다(첫 실사용 회귀)',
        not L.octave_suspect({'tail_st': -7.5, 'tail_peak_st': -5.0})
        and L.octave_suspect({'tail_st': -17.7, 'tail_peak_st': -15.0}))

    # K5 환각 거르기 — 가짜 받아쓰기로 transcribe() 전 구간
    bindir = os.path.join(tmp, 'bin_w')
    os.makedirs(bindir)
    fake_w = install_tool(bindir, 'whisper-cli', FAKE_WHISPER)
    # 0번 = 0.2초 «무성» 잡음 폭발 — 유성이면 R3 규칙상 버려도 누락이 된다(그건 R3·K5-c 가 따로 잰다)
    s = Synth(5).sil(0.5).sil(0.2, namp=3000).sil(0.8).tone(1.0).sil(0.8).tone(1.0).sil(0.5)
    x = s.arr()
    ch, _ = L.segment(L.frame_db(x))
    wenv = {'PATH': bindir, 'FAKE_WHISPER_TEXTS': json.dumps(['감사합니다.', '감사합니다.', '안녕하세요.'],
                                                             ensure_ascii=False)}
    wtmp = os.path.join(tmp, 'w5')
    os.makedirs(wtmp)
    words, dropped, missing, _q, _w = L.transcribe(x, ch, fake_w, '/dev/null', 'ko', wtmp, env=wenv)
    nmiss = len(missing) if isinstance(missing, list) else missing
    chk('K5 컨트롤 생존: 덩어리 3개 · 누락 0', len(ch) == 3 and nmiss == 0, 'chunks=%d missing=%s' % (len(ch), missing))
    chk('K5-a 0.2초 «감사합니다」 → 환각으로 버림', dropped == [0], str(dropped))
    chk('K5-b 1.0초 «감사합니다」 → 남김(진짜 끝인사일 수 있다)',
        any(w['chunk'] == 1 and '감사' in w['w'] for w in words), str([(w['chunk'], w['w']) for w in words]))

    # K5-c (codex R2-A): 빈 받아쓰기의 «누락」 판정은 길이가 아니라 유성음으로 — 짧은 실제 발화(0.3초 «네»)의
    #    빈 받아쓰기는 누락이고, 같은 길이의 무성 잡음 폭발(숨·마찰)의 빈 것은 누락이 아니다.
    s = Synth(6).sil(0.5).sil(0.3, namp=6000).sil(0.8).tone(0.3, amp=8000).sil(0.8).tone(1.0).sil(0.5)
    x = s.arr()
    ch, _ = L.segment(L.frame_db(x))
    wenv2 = dict(wenv, FAKE_WHISPER_TEXTS=json.dumps(['', '', '안녕하세요.'], ensure_ascii=False))
    wtmp2 = os.path.join(tmp, 'w5c')
    os.makedirs(wtmp2)
    _, _, miss_c, quiet_c, _w = L.transcribe(x, ch, fake_w, '/dev/null', 'ko', wtmp2, env=wenv2)
    chk('K5-c 컨트롤 생존: 덩어리 3개(무성 잡음 0.3초 · 유성 0.3초 · 1.0초)', len(ch) == 3, 'chunks=%d %s' % (len(ch), ch))
    chk('K5-c 0.3초 유성 덩어리의 빈 받아쓰기 = 누락 · 0.3초 무성 잡음의 빈 것은 누락 아님',
        list(miss_c) == [1] and list(quiet_c) == [0], 'missing=%s unvoiced_empty=%s' % (miss_c, quiet_c))

    # K6 «-는」 탈락 후보 — 받아쓰기에서 «는」 하나만 다르다
    script = [(1, '오늘은 놓치는 것을 고치는 법을 봅니다.'), (2, '저는 그렇게 봅니다.')]
    ch6 = [(0.0, 3.0), (3.5, 5.0)]

    def mk(t0, t1):
        out = [{'w': w, 'chunk': 0, 't0': 0, 't1': 1} for w in t0.split()]
        return out + [{'w': w, 'chunk': 1, 't0': 3.5, 't1': 4} for w in t1.split()]
    r_l = L.row_pron(script, mk('오늘은 놓친 것을 고치는 법을 봅니다.', '저는 그렇게 봅니다.'), ch6)
    r_k = L.row_pron(script, mk('오늘은 놓치는 것을 고치는 법을 봅니다.', '저는 그렇게 봅니다.'), ch6)
    chk('K6-a «놓치는→놓친」 → 관형형 탈락 1/2', [e['word'] for e in r_l['neun']['verb'][1]] == ['놓치는']
        and r_l['neun']['verb'][0] == 2, str(r_l['neun']))
    chk('K6-b 대명사 «저는」 은 따로 센다', r_l['neun']['pron'][0] == 1 and not r_l['neun']['pron'][1],
        str(r_l['neun']['pron']))
    chk('K6-c 같은 받아쓰기 → 탈락 0 · 치환 0', not r_k['neun']['verb'][1] and not r_k['subs'], str(r_k['neun']))

    # 퇴화 — 의존성 부재는 «측정불가」이고 exit 2, 0 이 아니다
    print('L16 퇴화 · CLI')
    wav = greet_fixture(os.path.join(tmp, 'rehearsal.wav'))
    wav_v = greet_fixture(os.path.join(tmp, 'rehearsal_voiced.wav'), blip='voiced')
    scr = os.path.join(tmp, 'script.txt')
    with open(scr, 'w', encoding='utf-8') as fh:
        fh.write(GREET_SCRIPT)
    empty = os.path.join(tmp, 'empty_bin')
    os.makedirs(empty)
    ffbin = os.path.join(tmp, 'bin_f')
    os.makedirs(ffbin)
    install_tool(ffbin, 'ffmpeg', FAKE_FFMPEG)
    install_tool(ffbin, 'ffprobe', FAKE_FFPROBE)
    both = os.path.join(tmp, 'bin_fw')
    os.makedirs(both)
    install_tool(both, 'ffmpeg', FAKE_FFMPEG)
    install_tool(both, 'ffprobe', FAKE_FFPROBE)
    install_tool(both, 'whisper-cli', FAKE_WHISPER)
    model = os.path.join(tmp, 'fake-model.bin')
    open(model, 'w').close()
    texts = json.dumps(GREET_TEXTS, ensure_ascii=False)

    r = L.coach(wav, scr, (), env={'PATH': empty})
    st = {x['n']: x['cause'] for x in r['rows']}
    chk('D1 ffmpeg 없음 → exit 2', r['rc'] == 2, 'rc=%d' % r['rc'])
    chk('D1-b ①~④ 전부 «측정불가(UNMEASURED)」, ⑤⑥ 은 «설계상 범위 밖」',
        all(st[n] == L.CAUSE_DEP for n in (1, 2, 3, 4)) and st[5] == st[6] == L.CAUSE_DESIGN, str(st))
    chk('D1-c 출력이 «통과 아님」을 말한다', '통과 아님' in L.render(r))

    real_ff = shutil.which('ffmpeg')
    if real_ff:
        r = L.coach(wav, None, {1, 2}, env={'PATH': os.path.dirname(real_ff), 'PREPREP_FFMPEG': '/nonexistent/ffmpeg'})
        chk('D2 명시 경로(PREPREP_FFMPEG)가 틀리면 PATH 로 몰래 갈아타지 않는다', r['rc'] == 2
            and 'PREPREP_FFMPEG' in r['deps']['ffmpeg'], r['deps']['ffmpeg'])
    else:
        skip('D2 명시 경로 우선', '이 기계에 진짜 ffmpeg 가 없다')

    r = L.coach(wav, scr, (), env={'PATH': ffbin})
    st = {x['n']: (x['status'], x['cause']) for x in r['rows']}
    chk('D3 whisper 없음 → ①② 측정불가(UNMEASURED)',
        st[1] == (L.ST_NONE, L.CAUSE_DEP) and st[2] == (L.ST_NONE, L.CAUSE_DEP), str(st))
    chk('D3-b whisper 없어도 ③④ 는 오디오만으로 돈다',
        st[3][1] is None and st[4][1] is None and r['n_chunks'] == 8, '%s chunks=%s' % (st, r.get('n_chunks')))
    chk('D3-c 그 상태의 종료코드는 2(통과 아님)', r['rc'] == 2, 'rc=%d' % r['rc'])

    r = L.coach(wav, scr, {1, 2}, env={'PATH': ffbin})
    chk('D4 --audio-only(선언) → exit 0, ①② 는 «범위에서 뺌」으로 남는다',
        r['rc'] == 0 and r['rows'][0]['cause'] == r['rows'][1]['cause'] == L.CAUSE_DECLARED, 'rc=%d' % r['rc'])

    r = L.coach(wav, scr, (), env={'PATH': both, 'FAKE_WHISPER_TEXTS': texts})
    chk('D5 모델 경로 미지정 → ①② 측정불가, 이유가 PREPREP_WHISPER_MODEL 을 댄다',
        r['rc'] == 2 and 'PREPREP_WHISPER_MODEL' in r['rows'][0]['why'], r['rows'][0]['why'])

    env = {'PATH': both, 'FAKE_WHISPER_TEXTS': texts, 'PREPREP_WHISPER_MODEL': model}
    r = L.coach(wav, scr, (), env=env)
    st = {x['n']: (x['status'], x['cause']) for x in r['rows']}
    chk('D6 전 의존성(가짜) → ①~④ 전부 측정 · exit 0', r['rc'] == 0 and all(st[n][1] is None for n in (1, 2, 3, 4)),
        'rc=%d %s' % (r['rc'], st))
    chk('D6-b 짧은 잡음의 «감사합니다」 1개를 버리고 그 수를 낸다', r.get('asr', {}).get('hallucination_dropped') == 1,
        str(r.get('asr')))
    chk('D6-c ① 은 «부분」(후보)이고 «놓치는」 탈락을 근거 시각과 함께 댄다',
        st[1][0] == L.ST_PARTIAL and any('놓치는' in d for _, d in r['rows'][0]['evidence']),
        str(r['rows'][0]['evidence']))
    chk('D6-d ② 가 홀로 선 «그런데」를 시각과 함께 댄다', any('그런데' in d for _, d in r['rows'][1]['evidence']),
        str(r['rows'][1]['evidence']))
    chk('D6-e ④ 작게 녹음한 끝인사가 가장 조용한 덩어리(1/8)로 잡힌다',
        any('끝 덩어리' in o and '조용한 쪽부터 1/8' in o for o in r['rows'][3]['obs']), str(r['rows'][3]['obs']))
    out = L.render(r)
    chk('D6-f 보정 상수가 «UNCALIBRATED」 로 찍힌다', out.count('UNCALIBRATED') >= 2)
    # B2(codex 교차 리뷰): 종전 단언은 상태값이 «L.ST_* 중 하나」인지만 봤다 — 같은 상수를 읽으니
    #    렌더러가 «통과」를 써도 초록이었다. 이제 «렌더된 글자」를 직접 본다.
    chk('B2 렌더 출력에 «통과」/PASS 낱말이 없다(부정문 «통과 아님」 제외)', not says_pass(out),
        '찾은 줄: %s' % [l for l in out.splitlines() if says_pass(l)][:2])
    orig = L.ST_MEASURED
    try:
        L.ST_MEASURED = '통과'            # 되돌림: 렌더러가 «통과」라고 말하게 만든다
        out_m = L.render(L.coach(wav, scr, (), env=env))
    finally:
        L.ST_MEASURED = orig
    chk('B2-b 되돌림: 상태명을 «통과」로 바꾸면 B2 단언이 빨개진다(장식 아님)', says_pass(out_m) and '통과' in out_m)

    # A1(codex 교차 리뷰): 덩어리 일부만 받아쓰기가 누락되면 노트만 달리고 ② 가 «잡음」· exit 0 이었다
    r = L.coach(wav, scr, {1}, env=dict(env, FAKE_WHISPER_SKIP='3'))
    r2 = r['rows'][1]
    chk('A1 덩어리 1/8 받아쓰기 누락(--skip 1) → ② 는 «부분」 · exit 2',
        r['rc'] == 2 and r2['status'] == L.ST_PARTIAL and '1/8' in r2['why'],
        'rc=%d ②=%s · %s' % (r['rc'], r2['status'], r2['why']))
    at = r.get('asr', {}).get('missing_at') or []
    chk('A1-b 누락 위치(시각)를 낸다', len(at) == 1 and at[0] in r2['why'], 'missing_at=%s' % at)
    # A2: 유효 JSON 인데 낱말 0 인 «말소리」 덩어리는 누락으로 센다.
    # 🟥 codex R2 로 뒤집었다 — 종전 단언은 «0.2초 덩어리의 빈 것은 안 센다」였고, 그 0.2초 덩어리는
    #    유성음(톤)이다. 길이만으로 «잡음」을 단정한 구멍을 시험이 «보존」하고 있었다.
    t2 = list(GREET_TEXTS)
    t2[3], t2[4] = '', ''              # 3 = 1.6초 말소리 · 4 = 0.2초 «유성» 덩어리(짧은 발화)
    r = L.coach(wav_v, scr, {1, 3}, env=dict(env, FAKE_WHISPER_TEXTS=json.dumps(t2, ensure_ascii=False)))
    chk('A2 말소리 덩어리 둘(1.6초 · 0.2초 유성)의 빈 받아쓰기 = 누락 2 · --skip 1,3 에서도 exit 2 · ② 일부 못 들음',
        r.get('asr', {}).get('missing') == 2 and r['rc'] == 2 and r['rows'][1]['cause'] == L.CAUSE_INCOMPLETE,
        'asr=%s rc=%d ②=%s' % (r.get('asr'), r['rc'], r['rows'][1]['cause']))
    t3 = list(GREET_TEXTS)
    t3[4] = ''                         # 짧은 유성 덩어리 «하나만» 비었다 — R2 가 재현한 그 모양
    r = L.coach(wav_v, scr, {1, 3}, env=dict(env, FAKE_WHISPER_TEXTS=json.dumps(t3, ensure_ascii=False)))
    chk('A2-b 0.2초 유성 덩어리 하나만 빈 받아쓰기(--skip 1,3) → exit 2 · ② «잡음」 아님',
        r['rc'] == 2 and r['rows'][1]['status'] != L.ST_MEASURED and r['rows'][1]['cause'] is not None,
        'rc=%d ②=%s/%s' % (r['rc'], r['rows'][1]['status'], r['rows'][1]['cause']))
    # R3 (codex R3): 같은 0.2초 «유성» 덩어리를 받아쓰기가 «감사합니다.」로 지어내면 환각 필터가 버린다 —
    #    그 버림이 누락 셈을 우회해 rc 0 이 났다(빈 받아쓰기면 rc 2 인데). 버린 덩어리도 유성이면 누락이다.
    r = L.coach(wav_v, scr, {1, 3}, env=env)          # GREET_TEXTS[4] = '감사합니다.'
    a_ = r.get('asr', {})
    chk('R3 0.2초 유성 덩어리의 «감사합니다」 환각 → 버리되 누락으로 센다 · --skip 1,3 에서 exit 2',
        a_.get('hallucination_dropped') == 1 and a_.get('missing') == 1 and r['rc'] == 2
        and r['rows'][1]['cause'] == L.CAUSE_INCOMPLETE, 'asr=%s rc=%d' % (a_, r['rc']))
    r = L.coach(wav, scr, {1, 3}, env=env)            # 같은 자리가 «무성» 잡음이면
    a_ = r.get('asr', {})
    chk('R3-b 컨트롤: 무성 잡음 덩어리의 «감사합니다」 환각 → 버리고 누락 아님 · exit 0',
        a_.get('hallucination_dropped') == 1 and a_.get('missing') == 0 and r['rc'] == 0, 'asr=%s rc=%d' % (a_, r['rc']))
    orig_vf = L.voiced_fraction
    try:
        L.voiced_fraction = lambda *a, **k: 0.0      # 되돌림: 유성 판정을 죽인다
        r = L.coach(wav_v, scr, {1, 3}, env=env)
    finally:
        L.voiced_fraction = orig_vf
    chk('R3-c 되돌림: 유성 판정을 죽이면 같은 오디오가 다시 exit 0 으로 접힌다(판정이 하중을 진다)',
        r['rc'] == 0 and r.get('asr', {}).get('missing') == 0, 'rc=%d asr=%s' % (r['rc'], r.get('asr')))
    # B-json (codex R2): 모양이 이상한 JSON 은 트레이스백이 아니라 그 덩어리의 누락이다
    try:
        r = L.coach(wav, scr, {1}, env=dict(env, FAKE_WHISPER_RAW=json.dumps(
            {'3': '[]', '5': '{"transcription": null}', '6': '{"transcription": [null]}'})))
        chk('B-json []·{"transcription":null}·[null] → 그 덩어리 누락 3 · exit 2',
            r.get('asr', {}).get('missing') == 3 and r['rc'] == 2, 'asr=%s rc=%d' % (r.get('asr'), r['rc']))
    except Exception as e:   # noqa: BLE001
        chk('B-json []·{"transcription":null}·[null] → 그 덩어리 누락 3 · exit 2', False,
            '예외로 죽음: %s: %s' % (type(e).__name__, e))
    # R4 (codex R4): 덩어리 하나의 이상한 «값」(거대 정수 · NaN · 음수 오프셋)이 전체를 죽이거나 조용히 통과했다.
    #    부류로 닫는다 — 덩어리 하나를 파싱하는 경계에서 어떤 예외든 그 덩어리의 누락(사유 보존)이다.
    def tok(frm, to):
        return '{"transcription":[{"tokens":[{"text":" 답이","offsets":{"from":%s,"to":%s},"p":0.9}]}]}' % (frm, to)
    raw4 = {'3': tok('1' + '0' * 400, '2' + '0' * 400), '5': tok('NaN', '100'), '6': tok('-500', '100')}
    try:
        r = L.coach(wav, scr, {1}, env=dict(env, FAKE_WHISPER_RAW=json.dumps(raw4)))
        a_ = r.get('asr', {})
        chk('R4 거대 정수 · NaN · 음수 오프셋 → 각 덩어리 누락 3 · exit 2 (크래시 아님 · 조용한 통과 아님)',
            a_.get('missing') == 3 and r['rc'] == 2, 'asr.missing=%s rc=%d' % (a_.get('missing'), r['rc']))
        why = a_.get('missing_why') or {}
        chk('R4-b 누락 사유 문자열을 덩어리마다 보존한다', len(why) == 3 and all(v for v in why.values()), str(why))
        chk('R4-b2 렌더 표에 못 들은 위치와 사유가 찍힌다', L.render(r).count('못 들음 ') >= 3)
    except Exception as e:   # noqa: BLE001 — 크래시 자체가 이 레인의 실패다
        chk('R4 거대 정수 · NaN · 음수 오프셋 → 각 덩어리 누락 3 · exit 2 (크래시 아님 · 조용한 통과 아님)', False,
            '예외로 죽음: %s' % type(e).__name__)
        chk('R4-b 누락 사유 문자열을 덩어리마다 보존한다', False, '위에서 죽음')
    orig_w = L.words_from_whisper
    try:
        def boom(*a, **k):
            raise RuntimeError('예상 못 한 파서 예외')
        L.words_from_whisper = boom       # 예외 «종류」 목록에 없는 것 — 부류로 닫혔나
        r = L.coach(wav, scr, {1}, env=env)
        chk('R4-c 목록에 없는 예외 종류(RuntimeError)도 그 덩어리 누락으로 갇힌다 · exit 2',
            r['rc'] == 2 and r['rows'][1]['cause'] == L.CAUSE_DEP
            and 'RuntimeError' in json.dumps(r.get('asr', {}), ensure_ascii=False),
            'rc=%d asr=%s' % (r['rc'], str(r.get('asr'))[:120]))
    except Exception as e:   # noqa: BLE001
        chk('R4-c 목록에 없는 예외 종류(RuntimeError)도 그 덩어리 누락으로 갇힌다 · exit 2', False,
            '예외로 죽음: %s' % type(e).__name__)
    finally:
        L.words_from_whisper = orig_w
    # ── 외부 입력 경계 전수(codex R5) — 경계마다 «임의 예외 주입」 1개. 각 경계는 그 경계 전용 예외로
    #    감싸여 «그 항목 UNMEASURED(사유) · exit 2」 로 떨어져야 하고, 크래시(exit 1)면 안 된다.
    print('L16 외부 입력 경계 — 임의 예외 주입')

    def inject(name, attr, exc, want_dep, skip_rows=(), script=scr, env_=None, mod=L):
        orig_ = getattr(mod, attr, None)
        had = hasattr(mod, attr)

        def raiser(*a, **k):
            raise exc('주입된 예외')
        setattr(mod, attr, raiser)
        try:
            r = L.coach(wav, script, skip_rows, env=env_ or env)
            dep = sorted(x['n'] for x in r['rows'] if x['cause'] == L.CAUSE_DEP)
            whys = ' '.join(x['why'] for x in r['rows'] if x['cause'] == L.CAUSE_DEP) + json.dumps(r['deps'], ensure_ascii=False)
            chk(name, r['rc'] == 2 and dep == sorted(want_dep) and exc.__name__ in whys,
                'rc=%d dep=%s why=%s' % (r['rc'], dep, whys[:90]))
        except Exception as e:   # noqa: BLE001 — 크래시 자체가 실패다
            chk(name, False, '예외로 죽음: %s: %s' % (type(e).__name__, e))
        finally:
            if had:
                setattr(mod, attr, orig_)
            else:
                delattr(mod, attr)

    inject('E1 환경변수 경로(도구 찾기) — which 가 터지면 ffmpeg 없음 → ①~④ UNMEASURED', 'which', ZeroDivisionError,
           [1, 2, 3, 4], mod=L.shutil)
    inject('E2 모델 경로 확인이 터지면 → ①② UNMEASURED, ③④ 는 돈다', 'check_model', ZeroDivisionError, [1, 2])
    inject('E3 녹화 디코드(ffmpeg 출력 포함)가 목록 밖 예외로 터지면 → ①~④ UNMEASURED', 'decode', ZeroDivisionError,
           [1, 2, 3, 4])
    inject('E4 원고 로드가 목록 밖 예외로 터지면 → ① 만 UNMEASURED', 'read_script', ZeroDivisionError, [1])
    inject('E5 whisper 실행이 목록 밖 예외로 터지면 → ①② UNMEASURED', 'run_whisper', ZeroDivisionError, [1, 2])
    inject('E6 덩어리 파일 준비(쓰기)가 터지면 → ①② UNMEASURED', 'write_wav', ZeroDivisionError, [1, 2])
    # E7 whisper 출력 파싱 경계는 R4-c(RuntimeError 주입)가 이미 잰다.
    try:
        import io
        from pptx import Presentation
        prs = Presentation()
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        ph = sl.notes_slide.notes_placeholder
        ph._element.getparent().remove(ph._element)       # 노트 본문 자리표시자 제거
        buf = io.BytesIO()
        prs.save(buf)
        pp = os.path.join(tmp, 'no_notes_placeholder.pptx')
        with open(pp, 'wb') as fh:
            fh.write(buf.getvalue())
        has_pptx = True
    except ImportError:
        has_pptx = False
    if has_pptx:
        try:
            r = L.coach(wav, pp, (), env=env)
            chk('E8 노트 자리표시자 없는 pptx → ① UNMEASURED · exit 2 (크래시 아님)',
                r['rc'] == 2 and r['rows'][0]['cause'] == L.CAUSE_DEP, 'rc=%d ①=%s' % (r['rc'], r['rows'][0]['why'][:60]))
        except Exception as e:   # noqa: BLE001
            chk('E8 노트 자리표시자 없는 pptx → ① UNMEASURED · exit 2 (크래시 아님)', False,
                '예외로 죽음: %s: %s' % (type(e).__name__, e))
    else:
        skip('E8 노트 자리표시자 없는 pptx', 'python-pptx 없음')
    # ── R6 (codex R6): 디코드가 rc 0 인데 손상 구간을 조용히 버리면 «다 들음」이 된다 ──
    print('L16 부분 디코드 — stderr 오류 · 길이 부족')

    def degraded(r):
        return (r['rc'] == 2 and all(r['rows'][n - 1]['cause'] == L.CAUSE_INCOMPLETE for n in (3, 4))
                and all(r['rows'][n - 1]['status'] == L.ST_PARTIAL for n in (3, 4)))
    msg = 'Decoding error: Invalid data found when processing input'
    for name, extra, needle in (
            ('G1 rc 0 + stderr 오류 줄 → ③④ «부분(디코드 오류)」 · exit 2', {'FAKE_FFMPEG_ERR': msg}, '디코드 오류 1줄'),
            ('G2 디코드 길이가 녹화보다 10% 짧음(stderr 없음) → ③④ «부분」 · exit 2', {'FAKE_FFMPEG_KEEP': '0.9'}, '짧음')):
        try:
            r = L.coach(wav, None, {1, 2}, env=dict(PATH=ffbin, **extra))
            chk(name, degraded(r) and needle in r['rows'][3]['why'],
                'rc=%d ④=%s/%s %s' % (r['rc'], r['rows'][3]['status'], r['rows'][3]['cause'], r['rows'][3]['why'][:80]))
        except Exception as e:   # noqa: BLE001
            chk(name, False, '예외로 죽음: %s: %s' % (type(e).__name__, e))
    r = L.coach(wav, None, {1, 2}, env={'PATH': ffbin, 'FAKE_FFMPEG_KEEP': '0.99'})
    chk('G3 컨트롤: 1% 짧음(문턱 2% 미만) · 오류 줄 없음 → exit 0 · 들린 표 그대로', r['rc'] == 0 and r['rows'][3]['cause'] is None,
        'rc=%d' % r['rc'])
    onlyff = os.path.join(tmp, 'bin_ffonly')
    os.makedirs(onlyff)
    install_tool(onlyff, 'ffmpeg', FAKE_FFMPEG)
    try:
        r = L.coach(wav, None, {1, 2}, env={'PATH': onlyff})
        chk('G4 ffprobe 없음 → 길이 대조 못 함 = «부분」 · exit 2 (못 잰 것을 다 들었다고 안 한다)',
            degraded(r) and '길이 대조' in r['rows'][3]['why'], 'rc=%d %s' % (r['rc'], r['rows'][3]['why'][:80]))
    except Exception as e:   # noqa: BLE001
        chk('G4 ffprobe 없음 → 길이 대조 못 함 = «부분」 · exit 2 (못 잰 것을 다 들었다고 안 한다)', False,
            '예외로 죽음: %s' % type(e).__name__)
    orig_p = getattr(L, 'probe_duration', None)
    try:
        def boomp(*a, **k):
            raise ZeroDivisionError('주입된 예외')
        L.probe_duration = boomp
        r = L.coach(wav, None, {1, 2}, env={'PATH': ffbin})
        chk('G5 길이 조회(ffprobe 출력) 경계에 예외 주입 → «부분」 · exit 2 (크래시 아님)',
            degraded(r) and 'ZeroDivisionError' in r['rows'][3]['why'], 'rc=%d' % r['rc'])
    except Exception as e:   # noqa: BLE001
        chk('G5 길이 조회(ffprobe 출력) 경계에 예외 주입 → «부분」 · exit 2 (크래시 아님)', False,
            '예외로 죽음: %s' % type(e).__name__)
    finally:
        if orig_p is None:
            del L.probe_duration
        else:
            L.probe_duration = orig_p
    # G6 진짜 ffmpeg 팔 — 합성 음성을 mp3 로 굽고 가운데 바이트를 깨뜨린다(codex 재현 형태: rc 0 + 오류 줄 + 잘림)
    if real_ff:
        mp3, bad3 = os.path.join(tmp, 'g6.mp3'), os.path.join(tmp, 'g6_bad.mp3')
        pr = subprocess.run([real_ff, '-nostdin', '-v', 'error', '-y', '-i', wav, '-c:a', 'libmp3lame', '-b:a', '64k', mp3],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if pr.returncode == 0 and os.path.isfile(mp3):
            d = bytearray(open(mp3, 'rb').read())
            for i in range(len(d) // 3, len(d) // 3 + 4000):
                d[i] = (d[i] * 7 + 13) & 255
            open(bad3, 'wb').write(bytes(d))
            renv = {'PATH': os.path.dirname(real_ff)}
            r_ok = L.coach(mp3, None, {1, 2}, env=renv)
            r_bad = L.coach(bad3, None, {1, 2}, env=renv)
            chk('G6 진짜 ffmpeg: 멀쩡한 mp3 → exit 0', r_ok['rc'] == 0, 'rc=%d %s' % (r_ok['rc'], r_ok.get('decode_issue')))
            chk('G6-b 진짜 ffmpeg: 가운데가 깨진 mp3(rc 0 디코드) → «부분」 · exit 2', degraded(r_bad),
                'rc=%d issue=%s' % (r_bad['rc'], str(r_bad.get('decode_issue'))[:100]))
        else:
            skip('G6 진짜 ffmpeg 손상 팔', 'libmp3lame 인코더 없음')
    else:
        skip('G6 진짜 ffmpeg 손상 팔', '이 기계에 ffmpeg 가 없다')
    # B-exec (codex R2): isfile·X_OK 는 통과하지만 실행이 안 되는 whisper(셰뱅 없는 스크립트 → ENOEXEC)
    noexec = os.path.join(tmp, 'bin_noexec')
    os.makedirs(noexec)
    install_tool(noexec, 'ffmpeg', FAKE_FFMPEG)
    install_tool(noexec, 'ffprobe', FAKE_FFPROBE)
    with open(os.path.join(noexec, 'whisper-cli'), 'wb') as fh:
        fh.write(b'\x00\x01not-an-executable\n')
    os.chmod(os.path.join(noexec, 'whisper-cli'), 0o755)
    try:
        r = L.coach(wav, scr, (), env={'PATH': noexec, 'PREPREP_WHISPER_MODEL': model})
        chk('B-exec 실행 불가 whisper → ①② 측정불가 · exit 2(트레이스백 아님)',
            r['rc'] == 2 and r['rows'][0]['cause'] == r['rows'][1]['cause'] == L.CAUSE_DEP
            and '실행' in r['rows'][1]['why'], 'rc=%d ②=%s' % (r['rc'], r['rows'][1]['why'][:70]))
    except Exception as e:   # noqa: BLE001
        chk('B-exec 실행 불가 whisper → ①② 측정불가 · exit 2(트레이스백 아님)', False,
            '예외로 죽음: %s: %s' % (type(e).__name__, e))
    # B1: 원고 인코딩 — CP949 는 읽고, 어느 쪽으로도 못 읽으면 ① 측정불가(트레이스백 아님)
    cp = os.path.join(tmp, 'script_cp949.txt')
    with open(cp, 'wb') as fh:
        fh.write(GREET_SCRIPT.encode('cp949'))
    bad = os.path.join(tmp, 'script_bad.txt')
    with open(bad, 'wb') as fh:
        fh.write(b'\xff\xff\xfe\xff' * 8)
    for label, path, want in (('B1 CP949 원고 → 읽는다(① 측정)', cp, 'ok'),
                              ('B1-b UTF-8·CP949 둘 다 아닌 원고 → ① 측정불가 · exit 2', bad, 'dep')):
        try:
            r = L.coach(wav, path, (), env=env)
            c1 = r['rows'][0]['cause']
            good = (c1 is None and r['rc'] == 0) if want == 'ok' else (c1 == L.CAUSE_DEP and r['rc'] == 2
                                                                        and '원고를 못 읽는다' in r['rows'][0]['why'])
            chk(label, good, 'rc=%d ①=%s' % (r['rc'], r['rows'][0]['why'][:60]))
        except Exception as e:   # noqa: BLE001 — 트레이스백 자체가 이 레인의 실패다
            chk(label, False, '예외로 죽음: %s: %s' % (type(e).__name__, e))
    if hasattr(os, 'geteuid') and os.geteuid() != 0:
        nor = os.path.join(tmp, 'script_noread.txt')
        with open(nor, 'w', encoding='utf-8') as fh:
            fh.write(GREET_SCRIPT)
        os.chmod(nor, 0)
        try:
            r = L.coach(wav, nor, (), env=env)
            chk('B1-c 읽기 권한 없는 원고 → ① 측정불가 · exit 2', r['rows'][0]['cause'] == L.CAUSE_DEP and r['rc'] == 2)
        except Exception as e:   # noqa: BLE001
            chk('B1-c 읽기 권한 없는 원고 → ① 측정불가 · exit 2', False, '예외로 죽음: %s' % type(e).__name__)
        finally:
            os.chmod(nor, 0o600)
    else:
        skip('B1-c 읽기 권한', 'root 로 돌아서 권한 거부를 재현 못 한다')

    r = L.coach(wav, scr, (), env=dict(env, FAKE_WHISPER_TEXTS=json.dumps([''] * 8)))
    chk('D6-g 빈 받아쓰기(낱말 0) → ①② 측정불가 · exit 2 («접속사 0개」로 접히지 않는다)',
        r['rc'] == 2 and r['rows'][1]['cause'] == L.CAUSE_DEP and r['rows'][0]['cause'] == L.CAUSE_DEP,
        'rc=%d ②=%s' % (r['rc'], r['rows'][1]['why']))

    p = subprocess.run([sys.executable, os.path.join(HERE, 'lane_speech_coach.py'), wav, '--audio-only'],
                       env=dict(os.environ, PATH=ffbin + os.pathsep + os.path.dirname(sys.executable)),
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    chk('D7 CLI --audio-only → rc 0 · 표 6행', p.returncode == 0 and p.stdout.decode().count(' — ') >= 6,
        'rc=%d' % p.returncode)
    p = subprocess.run([sys.executable, os.path.join(HERE, 'lane_speech_coach.py'), wav],
                       env=dict(os.environ, PATH=empty), stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    chk('D7-b CLI ffmpeg 없음 → rc 2 (1 도 0 도 아니다)', p.returncode == 2, 'rc=%d' % p.returncode)
    # codex R7: «²» 는 isdigit() 참·int() 실패 — 크래시(exit 1)가 아니라 인자 오류(exit 2)여야 한다.
    for bad in ('²', '9' * 30, '1,x'):
        p = subprocess.run([sys.executable, os.path.join(HERE, 'lane_speech_coach.py'), wav, '--skip', bad],
                           env=dict(os.environ, PATH=ffbin + os.pathsep + os.path.dirname(sys.executable)),
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        chk('D7-c CLI --skip %r → rc 2 (크래시 아님)' % bad,
            p.returncode == 2 and b'Traceback' not in p.stderr, 'rc=%d' % p.returncode)

    if real_ff:
        r = L.coach(wav, None, {1, 2}, env={'PATH': os.path.dirname(real_ff)})
        chk('D8 진짜 ffmpeg 팔 → 같은 덩어리 8개 · exit 0', r['rc'] == 0 and r['n_chunks'] == 8,
            'rc=%d chunks=%s' % (r['rc'], r.get('n_chunks')))
    else:
        skip('D8 진짜 ffmpeg 팔', '이 기계에 ffmpeg 가 없다')

    # 되돌림 — 지역 문턱을 전역 문턱으로 바꾸면 잡음 바닥이 오른 구간의 쉼이 사라져야 한다
    print('L16 되돌림 프로브')
    s = Synth(9)
    for _ in range(65):
        s.tone(0.6, amp=8000, namp=20, harmonics=1).sil(0.4, namp=20)
    for _ in range(65):
        s.tone(0.6, amp=8000, namp=400, harmonics=1).sil(0.4, namp=400)
    db = L.frame_db(s.arr())
    base, _ = L.segment(db)
    orig = L.thr_series
    try:
        L.thr_series = lambda d, *a, **k: [L.pct(d, 10) + L.C('rel_db')] * len(d)
        assert L.thr_series is not orig     # 적용 확인 — 바꿔 끼운 것이 실제로 쓰이는가
        mut, _ = L.segment(db)
    finally:
        L.thr_series = orig
    after, _ = L.segment(db)
    chk('R1 지역 문턱 → 잡음 바닥이 오른 구간도 덩어리로 가른다(≥90/130)', len(base) >= 90, 'BASE %d' % len(base))
    chk('R1-b 전역 문턱으로 되돌리면 그 구간이 붙는다(≥20 감소)', len(mut) <= len(base) - 20,
        'BASE %d → MUT %d' % (len(base), len(mut)))
    chk('R1-c 복원 후 BASE 와 같다', len(after) == len(base), 'AFTER %d' % len(after))


if __name__ == '__main__':
    sys.exit(main())
