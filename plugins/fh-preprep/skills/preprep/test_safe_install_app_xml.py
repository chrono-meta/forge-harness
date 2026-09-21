#!/usr/bin/env python3
"""test_safe_install_app_xml.py — `docProps/app.xml` 필수 부품 계약의 known-pair.

🟥 재는 것: **정본 v1.4 가 들어온 경로가 막혔나.**
   v1.4 는 `docProps/app.xml` 없이 설치돼 정본이 됐고, 그 결과
   ① 체크리스트 §F5(app.xml ↔ 파트 수 정합)가 검사 대상을 잃었고
   ② `ooxml/gate.py` 가 그 파트를 읽다 죽어 **한 줄도 검사 못 한 채 rc=1(발견)** 을 냈다.

   양팔은 **한 변수만 다르다** — 같은 덱에서 `docProps/app.xml` 을 뺐나 뒀나.
   컨트롤이 «볼 게 없어서» 통과하는 게 아님을 장 수 대조로 확인한다.

    python3 test_safe_install_app_xml.py      # 0 통과 / 1 실패
"""
import os, sys, io, shutil, zipfile, subprocess, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
SAFE = os.path.join(HERE, 'safe_install.py')

def make_deck(path, with_app=True):
    from pptx import Presentation
    tmp = path + '.src'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(tmp)
    with zipfile.ZipFile(tmp) as z:
        items = [(n, z.read(n)) for n in z.namelist()]
    if not with_app:
        items = [(n, d) for n, d in items if n != 'docProps/app.xml']
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, d in items: z.writestr(n, d)
    os.remove(tmp)
    with zipfile.ZipFile(path) as z:
        n_sl = len([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])
        has = 'docProps/app.xml' in z.namelist()
    return n_sl, has

def run(cand, dest, *extra):
    r = subprocess.run([sys.executable, SAFE, cand, dest, *extra], capture_output=True, text=True)
    return r.returncode, r.stdout + r.stderr

def main():
    tmp = tempfile.mkdtemp(prefix='safeinst_pair_')
    fails = []
    try:
        good = os.path.join(tmp, 'good.pptx');  n_g, h_g = make_deck(good, True)
        bad  = os.path.join(tmp, 'bad.pptx');   n_b, h_b = make_deck(bad,  False)
        print('safe_install.py · docProps/app.xml known-pair')
        print('  컨트롤 생존 확인: 장 수 %d == %d · app.xml 유무 %s vs %s' % (n_g, n_b, h_g, h_b))
        if n_g != n_b or h_g == h_b: fails.append('control-parity')

        # ARM A (known-positive): app.xml 없는 덱 → 거부(rc=1), 정본 안 생김
        destA = os.path.join(tmp, 'canonA.pptx')
        rcA, outA = run(bad, destA)
        okA = (rcA == 1 and not os.path.exists(destA))
        print('  %s A app.xml 없는 덱 → 거부      rc=%s · 정본 생성됨=%s'
              % ('✅' if okA else '❌', rcA, os.path.exists(destA)))
        if not okA: fails.append('A'); print('     ↳', ' / '.join(outA.strip().splitlines()[-2:]))
        if 'app.xml' not in outA:
            fails.append('A-wording'); print('  ❌ A 팔 메시지가 app.xml 을 이름으로 안 댄다')
        else:
            print('  ✅ A 팔이 원인을 이름으로 댄다(python-pptx 힌트 포함)')

        # ARM B (컨트롤): 같은 덱, app.xml 만 있음 → 통과(rc=0), 정본 생김
        destB = os.path.join(tmp, 'canonB.pptx')
        rcB, outB = run(good, destB)
        okB = (rcB == 0 and os.path.exists(destB))
        print('  %s B app.xml 있는 덱 → 통과      rc=%s · 정본 생성됨=%s'
              % ('✅' if okB else '❌', rcB, os.path.exists(destB)))
        if not okB: fails.append('B'); print('     ↳', ' / '.join(outB.strip().splitlines()[-3:]))

        print('PASS' if not fails else 'FAIL · ' + ', '.join(fails))
        return 1 if fails else 0
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

if __name__ == '__main__':
    sys.exit(main())
