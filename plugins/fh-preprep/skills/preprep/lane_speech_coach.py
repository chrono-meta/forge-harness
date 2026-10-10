#!/usr/bin/env python3
"""lane_speech_coach.py — L16 speech-coach: 리허설 녹화를 듣고 코칭 6항 표를 낸다. (advisory)

다른 레인은 «표면 사이»(덱 ↔ 원고 ↔ 노트)를 본다. 이 레인은 그 원고를 **소리 내 읽은 결과**를
본다 — 발표 준비의 마지막 표면이다. 입력이 `surfaces.yaml` 의 문서가 아니라 녹화라서
`preprep.py` 에 배선하지 않고 독립 CLI 로 부른다.

    python3 lane_speech_coach.py <녹화.mp4|.wav> [--script 덱.pptx|원고.txt] [--audio-only]
                                 [--skip 1,2] [--json 결과.json]

코칭 6항(아나운서 코칭 한 회분을 known-positive 로 삼아 나눈 축):
    ① 소리 내 읽기·발음      — 받아쓰기 × 원고 대조. 어긋난 자리는 «발음 후보»일 뿐이다
    ② 쉼표 여유 · 접속사 끌기 — 쉼으로 자른 덩어리 × 받아쓰기. 접속사 앞뒤 쉼과 늘임
    ③ 어미 밀당              — 문장 끝 300ms 피치 모양(하강 · 올림 · 올렸다 내림)
    ④ 첫 인사·끝인사 강조     — 첫·끝 덩어리의 크기를 녹화 자기 분포에 대고 잰다
    ⑤ 시선                   — 오디오 계기 범위 밖(설계상)
    ⑥ 스피치보다 내용         — 계기 대상이 아니다. 이 스킬의 다른 레인(L1~L15)이 맡는다

의존성 — 전부 로컬, 녹화는 기계 밖으로 나가지 않는다:
    ffmpeg       필수.  없으면 «전부 측정불가» · exit 2.  경로 지정: PREPREP_FFMPEG
    whisper-cli  선택.  없으면 ①② 측정불가, ③④ 는 오디오만으로 돈다.
                 경로 지정: PREPREP_WHISPER_CLI · 모델: PREPREP_WHISPER_MODEL(기본값 없음 — 지정 필수)
                 언어: PREPREP_WHISPER_LANG(기본 ko)
    python-pptx  `--script *.pptx` 를 줄 때만(발표자 노트를 원고로 읽는다)

종료코드 — 🟥 **1 은 쓰지 않는다.**
    0  범위 안의 항목이 전부 측정됐다. 🟥 «좋은 발표」가 아니라 «다 들었다」다
    2  하나라도 측정 못 했다(ffmpeg 없음 · 디코드 실패 · 발화 0 · whisper/모델/원고 없음).
       범위에서 빼려면 `--audio-only`(①② 제외) 또는 `--skip` 으로 «선언»한다 — 선언한 것은
       «범위에서 뺌(선언)」으로 출력된다. 조용히 사라지지 않는다.
    이 레인의 숫자는 묘사다. 문턱(끌기 · 밀당 · 덩어리 경계)은 녹화 1건 · TTS 대조군에서 고른
    값이라 전부 `UNCALIBRATED` 로 출력된다. 숫자를 «발견」으로 바꿔 종료코드에 태우면 그건
    보정 안 된 계기가 판정을 내리는 것이다 — L8·R*·P* 와 같은 advisory 관례.

알려진 한계(인용 전에 읽어라):
    · 받아쓰기 낱말 시각은 실측 최대 ~0.8초 어긋났다. 그래서 낱말 시각으로 쉼을 재지 않는다 —
      쉼으로 자른 덩어리를 하나씩 받아쓰게 하고, 경계는 음향(덩어리)에서 가져온다.
      덩어리 안 낱말의 길이는 토큰 시각이라 `UNCALIBRATED` 로 표기한다.
    · 덩어리 단독 받아쓰기는 0.35초 미만 잡음에 «감사합니다」를 지어낸다(실측 16/550). 걸러내고
      그 수를 출력한다. 다른 형태의 환각은 못 거른다.
    · ① 의 어긋남은 받아쓰기 오류와 실제 발음을 이 계기만으로 못 가른다 — 사람 귀 확인 전에는
      «후보」다.
    · ③ 은 기준선이 없다. 한국어 평서문은 원래 내려가므로 «하강 비율이 높다」는 결함 판정이 아니다.
    · 피치는 자기상관 근사다(순음·합성 톱니 known-pair 통과, 실화자 프레임 단위 손검증 없음).
"""
import argparse
import array
import difflib
import json
import math
import operator
import os
import re
import shutil
import subprocess
import sys
import tempfile
import wave

SR = 16000
HOP = 160            # 10ms
LANE = 'L16 speech-coach'

# 🟥 보정 상수 — 전부 UNCALIBRATED. 녹화 1건(화자 1명) · TTS 대조군에서 고른 값이다.
#    출력에 그대로 싣는다. 이 값을 근거로 «빠르다/단조롭다/끌지 않았다」를 판정하지 마라.
CAL = {
    'rel_db':           (12.0, '발화 문턱 = 지역 하위 10% + 이 값(dB)'),
    'thr_win_s':        (60,   '지역 문턱 창(초) — 잡음 바닥이 구간마다 다르다(실측 ~10dB 차)'),
    'thr_step_s':       (10,   '지역 문턱 갱신 간격(초)'),
    'min_sil_s':        (0.15, '이보다 짧은 무음은 덩어리 경계로 안 본다'),
    'min_speech_s':     (0.08, '이보다 짧은 «발화」는 무음으로 흡수'),
    'pad_s':            (0.12, '덩어리를 받아쓰기에 넘길 때 앞뒤 여유'),
    'hallu_max_s':      (0.35, '이보다 짧은 덩어리의 상투 인사 받아쓰기는 환각으로 버린다(유성이면 누락으로도 센다)'),
    'voiced_min':       (0.25, '빈 받아쓰기 덩어리의 유성 비율이 이 이상이면 «말소리를 놓쳤다」(누락)로 센다'),
    'drag_stretch':     (1.30, '접속사 «끌기」 후보 문턱(TTS 보통 1.11 · 끈 것 1.44 사이)'),
    'phrase_end_gap_s': (0.60, '받아쓰기 없을 때 «구 끝」 대리: 뒤 쉼이 이 이상'),
    'tail_s':           (0.30, '문장 끝 창(초)'),
    'body_s':           (1.20, '문장 끝 직전 본문 창(초)'),
    'peak_st':          (1.0,  '밀당 = 끝 창 봉우리가 본문 대비 이 반음 이상 올랐다가 그만큼 내려옴'),
    'decode_short_frac': (0.02, '디코드된 길이가 녹화 길이(ffprobe)보다 이 비율 이상 짧으면 «부분 디코드」'),
    'octave_up_st':     (6.0,  '끝 봉우리가 본문 대비 이 반음 넘게 오르면 옥타브 오류 의심 — 예시·밀당에서 뺀다'),
    'octave_down_st':   (12.0, '끝 평균이 본문 대비 이 반음 넘게 내리면 옥타브 오류 의심(한국어 평서 하강은 보통 이 안쪽)'),
}


def C(k):
    return CAL[k][0]


CONJ = ['그런데', '그리고', '그래서', '그러나', '그러면', '하지만', '그래도', '그러니까', '따라서', '반면', '또한', '즉']
STOCK_HALLU = re.compile(r'^(감사합니다|고맙습니다|시청해주셔서감사합니다)[.!]?$')
PRONOUN_TOPIC = ('저', '우리', '이', '그', '이것', '그것', '나', '너')

ST_MEASURED, ST_PARTIAL, ST_NONE = '잡음', '부분', '측정불가'
CAUSE_DEP, CAUSE_DESIGN, CAUSE_DECLARED = 'unmeasured', 'by_design', 'declared'
CAUSE_INCOMPLETE = 'incomplete'   # 측정은 됐지만 일부 덩어리를 못 들었다 — rc 에 센다


# ── 오디오 기초 ───────────────────────────────────────────────────────────────────────────────
def fmt_t(s):
    if s is None:
        return '—'
    m = int(s // 60)
    return '%d:%04.1f' % (m, s - 60 * m)


def pct(v, p):
    s = sorted(v)
    if not s:
        return float('nan')
    return s[min(len(s) - 1, int(p / 100 * len(s)))]


def frame_db(x):
    """10ms 프레임 RMS dBFS."""
    mul = operator.mul
    out = []
    for i in range(0, len(x) - HOP + 1, HOP):
        w = x[i:i + HOP]
        out.append(20 * math.log10(math.sqrt(sum(map(mul, w, w)) / HOP) + 1e-9) - 90.309)
    return out


def thr_series(db, win_s=None, step_s=None, rel=None):
    """지역 적응 문턱: step 마다 ±win/2 창의 하위 10% + rel dB.

    전역 문턱 하나로는 잡음 바닥이 올라간 구간의 쉼이 «발화」로 읽혀 덩어리가 수십 초로 붙는다
    (실측 21초짜리 덩어리). 되돌림 프로브가 이 함수를 전역 문턱으로 바꿔 그 병합을 재현한다.
    """
    win_s = C('thr_win_s') if win_s is None else win_s
    step_s = C('thr_step_s') if step_s is None else step_s
    rel = C('rel_db') if rel is None else rel
    n = len(db)
    out = [0.0] * n
    step = int(step_s * 100)
    half = int(win_s * 50)
    for c in range(0, n, step):
        w = db[max(0, c + step // 2 - half): c + step // 2 + half]
        t = pct(w, 10) + rel
        for i in range(c, min(n, c + step)):
            out[i] = t
    return out


def runs(db, thr, min_sil=None):
    """(발화?, 시작초, 끝초) 런. min_sil 미만 무음은 앞 발화에 흡수."""
    min_sil = C('min_sil_s') if min_sil is None else min_sil
    if not db:
        return []
    v = [d > t for d, t in zip(db, thr)]
    rr, cur, st = [], v[0], 0
    for i in range(1, len(v)):
        if v[i] != cur:
            rr.append([cur, st, i])
            cur, st = v[i], i
    rr.append([cur, st, len(v)])
    m = []
    for r in rr:
        if not r[0] and (r[2] - r[1]) * 0.01 < min_sil and m:
            m[-1][2] = r[2]
            continue
        if m and m[-1][0] == r[0]:
            m[-1][2] = r[2]
        else:
            m.append(r)
    return [(bool(r[0]), r[1] * 0.01, r[2] * 0.01) for r in m]


def clean_runs(R, min_speech=None):
    min_speech = C('min_speech_s') if min_speech is None else min_speech
    out = []
    for v, a, b in R:
        if v and b - a < min_speech:
            v = False
        if out and out[-1][0] == v:
            out[-1] = (v, out[-1][1], b)
        else:
            out.append((v, a, b))
    return out


def segment(db):
    """발화 덩어리 [(a,b)] · 문턱 시계열. 쉼으로 갈린 덩어리가 이 레인의 시간 정본이다."""
    thr = thr_series(db)
    R = clean_runs(runs(db, thr))
    return [(a, b) for v, a, b in R if v], thr


def smooth(v, k=5):
    h = k // 2
    out = []
    for i in range(len(v)):
        w = v[max(0, i - h): i + h + 1]
        out.append(sum(w) / len(w))
    return out


def peaks_in(db, thr_v, f0, f1):
    """f0..f1 프레임 구간의 음절 핵(평활 에너지 봉우리, 돌출 3dB · 간격 80ms)."""
    lo_ = max(0, f0 - 20)
    seg = smooth(db[lo_:f1 + 20], 5)
    out, last = [], -999
    for i in range(1, len(seg) - 1):
        g = lo_ + i
        if g < f0 or g >= f1:
            continue
        if seg[i] > thr_v and seg[i] >= seg[i - 1] and seg[i] > seg[i + 1]:
            lo = min(seg[max(0, i - 15):i] or [seg[i]])
            ro = min(seg[i + 1:i + 16] or [seg[i]])
            if seg[i] - max(lo, ro) >= 3.0 and g - last >= 8:
                out.append(g)
                last = g
    return out


def f0_at(x, center_s):
    """40ms 창 · 4kHz 자기상관 피치(50~400Hz). 서브하모닉 회피: 최댓값 0.85 이상인 첫 국소 최대."""
    c = int(center_s * SR)
    raw = x[c:c + (160 + 80) * 4]
    if len(raw) < (160 + 80) * 4:
        return None
    d = [(raw[i] + raw[i + 1] + raw[i + 2] + raw[i + 3]) / 4.0 for i in range(0, len(raw) - 3, 4)]
    W = 160
    m = sum(d[:W]) / W
    d = [v - m for v in d]
    a = d[:W]
    mul = operator.mul
    e0 = sum(map(mul, a, a)) + 1e-9
    rs = [0.0] * 81
    for lag in range(10, 81):
        b = d[lag:lag + W]
        rs[lag] = sum(map(mul, a, b)) / math.sqrt(e0 * (sum(map(mul, b, b)) + 1e-9))
    best = max(rs[10:81])
    bl = None
    for lag in range(11, 80):
        if rs[lag] >= 0.85 * best and rs[lag] >= rs[lag - 1] and rs[lag] >= rs[lag + 1]:
            bl = lag
            break
    return 4000.0 / bl if (bl and best > 0.55) else None


def contour(x, db, thr, t0, t1, step=0.02):
    pts = []
    t = t0
    while t < t1:
        fi = int(t * 100)
        if 0 <= fi < len(db) and db[fi] > thr[fi] + 3:
            pts.append((round(t, 2), f0_at(x, t)))
        else:
            pts.append((round(t, 2), None))
        t += step
    idx = [i for i, (_, v) in enumerate(pts) if v]
    clean = list(pts)
    for k, i in enumerate(idx):   # 옥타브 오류: 이웃 7개 중앙값 ×1.6 / ×0.62 밖은 버린다
        nb = sorted(pts[j][1] for j in idx[max(0, k - 3):k + 4])
        med = nb[len(nb) // 2]
        if pts[i][1] > med * 1.6 or pts[i][1] < med * 0.62:
            clean[i] = (pts[i][0], None)
    return clean


def st(a, ref):
    return 12 * math.log2(a / ref)


def end_metrics(x, db, thr, ra, rb):
    """문장(구) 끝 rb 직전: 끝 창 평균 피치(본문 중앙값 대비 반음) · 봉우리 · 마지막 0.1초 · 에너지 낙차."""
    tail_s, body_s = C('tail_s'), C('body_s')
    t0 = max(ra, rb - body_s)
    v = [(t, f) for t, f in contour(x, db, thr, t0, rb) if f]
    tail = [f for t, f in v if t >= rb - tail_s]
    body = [f for t, f in v if t < rb - tail_s]
    rec = {'end': round(rb, 2), 'n_voiced': len(v)}
    if len(tail) >= 3 and len(body) >= 3:
        ref = sorted(body)[len(body) // 2]
        rec['tail_st'] = round(st(sum(tail) / len(tail), ref), 1)
        rec['tail_peak_st'] = round(st(max(tail), ref), 1)
        last = [f for t, f in v if t >= rb - 0.1]
        if last:
            rec['tail_last_st'] = round(st(sum(last) / len(last), ref), 1)
    e_tail = db[int((rb - tail_s) * 100):int(rb * 100)]
    e_body = db[int(t0 * 100):int((rb - tail_s) * 100)]
    if e_tail and e_body:
        rec['tail_energy_drop_db'] = round(pct(e_tail, 50) - pct(e_body, 50), 1)
    return rec


def is_push_pull(rec):
    """밀당 = 끝 창에서 한 번 올렸다(≥peak_st) 내려옴(봉우리 대비 ≥peak_st)."""
    p, l = rec.get('tail_peak_st'), rec.get('tail_last_st')
    return (p is not None and l is not None and p >= C('peak_st') and p - l >= C('peak_st')
            and not octave_suspect(rec))


def octave_suspect(rec):
    """봉우리가 +octave_up_st 넘게 오르거나 끝 평균이 −octave_down_st 넘게 내림.

    🟥 위아래가 비대칭인 이유(첫 실사용에서 드러남): 처음엔 ±6 하나로 걸었는데, 한국어 평서문 끝은
       중앙 −7.5반음으로 «정상적으로» 깊게 내려가서 실녹화 214개 중 164개가 «의심」으로 빠졌다.
       위로 6반음 넘는 봉우리(실측 +13.2)는 드물고 옥타브 오류일 가능성이 크지만, 아래로는 한 옥타브
       가까이 가야 의심할 만하다.
    """
    return ((rec.get('tail_peak_st') or 0) > C('octave_up_st')
            or (rec.get('tail_st') or 0) < -C('octave_down_st'))


# ── 외부 도구 (로컬) ──────────────────────────────────────────────────────────────────────────
def _exc(e):
    return '%s: %s' % (type(e).__name__, str(e)[:120])


# ── 외부 입력 경계 ───────────────────────────────────────────────────────────────────────────
# 🟥 «외부 입력 하나가 전체를 죽인다」가 세 번 나왔다(codex R2 B-json · R4 거대 정수 · R5 노트 자리표시자).
#    사례별로 막지 않고 경계를 전수로 세운다. 각 경계는 «그 경계 전용 예외 → 그 항목 UNMEASURED(사유) ·
#    exit 2」로 떨어진다. 경계 밖(우리 분석 코드)은 넓게 잡지 않는다 — 우리 버그를 «못 들음」으로 숨기지 않게.
#      E1 환경변수 경로(도구 찾기)   find_tool          → ffmpeg: ①~④ / whisper: ①②
#      E2 모델 경로                  check_model        → ①②
#      E3 녹화 디코드(ffmpeg 출력)    coach: decode 호출  → ①~④ (rc 0 + stderr 오류 = «부분」 강등, codex R6)
#      E3b 녹화 길이(ffprobe 출력)    coach: probe 호출   → 대조 못 하면 ①~④ «부분」
#      E4 원고 로드(pptx/txt/md)     coach: read_script  → ①
#      E5 whisper 실행               run_whisper         → ①②
#      E6 덩어리 파일 준비            transcribe: 쓰기     → ①②
#      E7 whisper 출력(덩어리별)      transcribe: 파싱     → 그 덩어리 누락 → ①② 부분 · exit 2
def find_tool(env, envvar, name):
    """명시 경로가 있으면 그것만 쓴다(없으면 «없음」 — PATH 로 몰래 갈아타지 않는다). [경계 E1]"""
    try:
        p = env.get(envvar)
        if p:
            return (p, None) if (os.path.isfile(p) and os.access(p, os.X_OK)) else \
                   (None, '%s=%s 가 실행 파일이 아니다' % (envvar, p))
        w = shutil.which(name, path=env.get('PATH', ''))
        return (w, None) if w else (None, '%s 가 PATH 에 없다(%s 로 지정 가능)' % (name, envvar))
    except Exception as e:   # noqa: BLE001 — 경계 E1
        return None, '%s 경로 확인 실패(%s)' % (name, _exc(e))


def check_model(path):
    """[경계 E2] 모델 파일이 있나. 없으면 사유 문자열, 있으면 None."""
    return None if os.path.isfile(path) else 'PREPREP_WHISPER_MODEL 이 가리키는 파일이 없다'


def run_whisper(cmd, env):
    """[경계 E5] whisper-cli 한 번 실행."""
    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=env)


def local_recording_input(path):
    """Treat recording names as local files, never as options or protocol URLs."""
    return 'file:' + os.path.abspath(os.fspath(path))


def decode(path, ffmpeg, env=None):
    """반환 (샘플, stderr 오류 줄들).

    🟥 rc 0 은 «다 디코드했다」가 아니다(codex R6, 진짜 ffmpeg 로 재현): 손상 패킷을 만나면 ffmpeg 는
       `-v error` 로 «Invalid data found…」를 찍고 그 구간을 버린 채 rc 0 으로 끝난다. 그래서 stderr 를
       버리지 않고 돌려준다 — 호출자가 «부분 디코드」로 강등한다.
    """
    p = subprocess.run([ffmpeg, '-nostdin', '-v', 'error', '-i', local_recording_input(path), '-vn', '-ac', '1',
                        '-ar', str(SR), '-f', 's16le', '-'],
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
    if p.returncode != 0:
        raise RuntimeError('ffmpeg rc=%d: %s' % (p.returncode, p.stderr.decode('utf-8', 'replace')[-300:]))
    raw = p.stdout[:len(p.stdout) - len(p.stdout) % 2]
    a = array.array('h')
    a.frombytes(raw)
    if sys.byteorder == 'big':
        a.byteswap()
    errs = [l.strip() for l in p.stderr.decode('utf-8', 'replace').splitlines() if l.strip()]
    return a, errs


def probe_duration(path, ffprobe, env=None):
    """[경계 E3b] 녹화의 오디오 길이(초) — 오디오 스트림 길이, 없으면 컨테이너 길이."""
    p = subprocess.run([ffprobe, '-v', 'error', '-select_streams', 'a:0', '-show_entries',
                        'stream=duration:format=duration', '-of', 'default=nw=1',
                        '-i', local_recording_input(path)],
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE, env=env)
    if p.returncode != 0:
        raise RuntimeError('ffprobe rc=%d' % p.returncode)
    for l in p.stdout.decode('utf-8', 'replace').splitlines():
        v = l.split('=', 1)[-1].strip()
        try:
            d = float(v)
        except ValueError:
            continue           # N/A
        if math.isfinite(d) and d > 0:
            return d
    raise ValueError('길이를 못 읽었다')


def write_wav(path, samples):
    with wave.open(path, 'wb') as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(SR)
        s = array.array('h', samples)
        if sys.byteorder == 'big':
            s.byteswap()
        w.writeframes(s.tobytes())


def words_from_whisper(tj, offset):
    """whisper-cli -ojf JSON → 낱말 [{w,t0,t1,p}]. 시각은 덩어리 파일 기준 + offset."""
    if not isinstance(tj, dict) or not isinstance(tj.get('transcription'), list):
        raise ValueError('받아쓰기 JSON 모양이 아니다')
    W = []
    for s in tj['transcription']:
        if not isinstance(s, dict) or not isinstance(s.get('tokens', []), list):
            raise ValueError('받아쓰기 구간 모양이 아니다')
        cur = None
        for tk in s.get('tokens', []):
            t = tk.get('text', '')
            if t.startswith('[_') or t.startswith('<|'):
                continue
            fr, to = tk['offsets']['from'], tk['offsets']['to']
            # 값 검증 — 모양만 맞고 값이 터무니없는 것(거대 정수 · NaN · 음수 · 뒤집힘)도 «못 읽음」이다
            for v in (fr, to):
                if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v < 0:
                    raise ValueError('토큰 오프셋 값이 이상하다: %r' % (v,))
            if to < fr:
                raise ValueError('토큰 오프셋이 뒤집혔다: %r > %r' % (fr, to))
            a = fr / 1000 + offset
            b = to / 1000 + offset
            if t.startswith(' ') or cur is None:
                if cur:
                    W.append(cur)
                cur = {'w': t.strip(), 't0': round(a, 2), 't1': round(b, 2), 'p': tk.get('p', 1.0)}
            else:
                cur['w'] += t
                cur['t1'] = round(b, 2)
                cur['p'] = min(cur['p'], tk.get('p', 1.0))
        if cur:
            W.append(cur)
    return [w for w in W if w['w']]


def voiced_fraction(x, a, b, step=0.02):
    """덩어리 a..b 의 유성 비율(자기상관 피치가 잡힌 20ms 점의 몫). 무성 잡음(숨·마찰·클릭)은 낮다."""
    ts = []
    t = a
    while t < b:
        ts.append(t)
        t += step
    if not ts:
        return 0.0
    return sum(1 for t in ts if f0_at(x, t)) / len(ts)


class WhisperFailed(Exception):
    pass


def transcribe(x, chunks, whisper, model, lang, tmp, batch=100, env=None):
    """덩어리마다 따로 받아쓰기. 반환 (낱말들, 버린 환각 덩어리, 누락 덩어리 번호들, 무성이라 뺀 빈 덩어리,
    누락 사유 {덩어리: 사유}).

    누락 = JSON 이 없음 · 못 읽음 · 모양이 이상함 · **유효한데 낱말 0 인 말소리 덩어리**.
    🟥 마지막 것을 안 세면 «소리는 있는데 받아쓰기가 비었다」가 커버리지 구멍으로 조용히 빠진다
       (codex 교차 리뷰 A2). 🟥 «말소리」 판정은 **길이가 아니라 유성 비율**이다(codex R2) — 처음엔
       0.35초 미만을 잡음으로 단정했는데, 0.3초 «즉」·«네」 같은 실제 발화의 빈 받아쓰기가 누락에도
       환각에도 안 잡혀 exit 0 이 났다. 유성 비율이 voiced_min 미만인 무성 잡음만 누락에서 뺀다.
    whisper 를 아예 실행하지 못하면(ENOEXEC · 권한 …) WhisperFailed — ①② 측정불가로 올린다.
    """
    pad = C('pad_s')
    files = []
    try:                                       # 경계 E6 — 덩어리 파일 준비
        for j, (a, b) in enumerate(chunks):
            s0 = max(0, int((a - pad) * SR))
            s1 = min(len(x), int((b + pad) * SR))
            p = os.path.join(tmp, 'c%04d.wav' % j)
            write_wav(p, x[s0:s1])
            files.append((p, s0 / SR))
    except Exception as e:   # noqa: BLE001 — 경계 E6
        raise WhisperFailed('덩어리 파일 준비 실패(%s)' % _exc(e))
    for i in range(0, len(files), batch):
        try:                                   # 경계 E5 — isfile·X_OK 를 통과해도 실행이 안 될 수 있다(ENOEXEC 등)
            run_whisper([whisper, '-m', model, '-l', lang, '-ojf', '-np'] + [f for f, _ in files[i:i + batch]], env)
        except Exception as e:   # noqa: BLE001 — 경계 E5
            raise WhisperFailed('whisper-cli 실행 불가(%s)' % (
                e.strerror if isinstance(e, OSError) and e.strerror else _exc(e)))
    words, dropped, missing, quiet_empty, why = [], [], [], [], {}
    for j, (p, off) in enumerate(files):
        jp = p + '.json'
        if not os.path.exists(jp):
            missing.append(j)
            why[j] = 'JSON 없음'
            continue
        # 🟥 [경계 E7] 덩어리 «하나」를 파싱하는 경계 — 여기서는 어떤 예외든 그 덩어리의 누락이다(사유 보존).
        #    예외 종류를 나열해 막으면 다음 종류가 또 전체를 죽인다(codex R2 B-json → R4 OverflowError,
        #    같은 부류 두 번). 경계는 «읽기 + 파싱」 두 줄뿐이다 — 그 밖(전체 흐름)은 넓게 잡지 않는다.
        try:
            with open(jp, encoding='utf-8') as fh:
                ws = words_from_whisper(json.load(fh), off)
        except Exception as e:   # noqa: BLE001 — 의도된 부류 경계
            missing.append(j)
            why[j] = '%s: %s' % (type(e).__name__, str(e)[:120])
            continue
        a, b = chunks[j]
        if not ws:
            if voiced_fraction(x, a, b) >= C('voiced_min'):
                missing.append(j)   # 말소리가 있는데 받아쓰기가 비었다
                why[j] = '유성 덩어리의 빈 받아쓰기'
            else:
                quiet_empty.append(j)
            continue
        joined = re.sub(r'\s+', '', ''.join(w['w'] for w in ws))
        if b - a < C('hallu_max_s') and STOCK_HALLU.match(joined):
            dropped.append(j)
            # 🟥 버린 덩어리도 유성이면 «말은 했는데 못 알아들었다」 = 누락이다(codex R3). 이 줄이 없으면
            #    환각 필터가 위의 빈-받아쓰기 커버리지 검사를 우회한다 — 같은 오디오가 빈 받아쓰기면
            #    exit 2, «감사합니다」로 지어내면 exit 0 이 됐다. 무성이면 잡음으로 보고 누락에 안 센다.
            if voiced_fraction(x, a, b) >= C('voiced_min'):
                missing.append(j)
                why[j] = '유성 덩어리의 상투 문구 환각(못 알아들음)'
            continue
        for w in ws:
            w['chunk'] = j
            words.append(w)
    return words, dropped, missing, quiet_empty, why


# ── 원고 ─────────────────────────────────────────────────────────────────────────────────────
class ScriptUnreadable(Exception):
    pass


def read_script(path):
    """[(단위, 줄)]. 단위 = pptx 의 장 번호 또는 텍스트의 줄 번호. 무대 지시([간지]/[진행])는 뺀다."""
    if not os.path.isfile(path):
        raise ScriptUnreadable('원고 파일이 없다')
    ext = os.path.splitext(path)[1].lower()
    units = []
    if ext == '.pptx':
        try:
            from pptx import Presentation
        except ImportError:
            raise ScriptUnreadable('python-pptx 없음 — pptx 노트를 못 읽는다')
        try:
            prs = Presentation(path)
        except Exception as e:   # noqa: BLE001 — 깨진 파일은 «못 읽음」으로 계상한다
            raise ScriptUnreadable('pptx 를 못 연다: %s' % type(e).__name__)
        for i, s in enumerate(prs.slides, 1):
            if s.has_notes_slide:
                tf = s.notes_slide.notes_text_frame   # 노트 본문 자리표시자가 없으면 None(codex R5)
                if tf is not None:
                    units.append((i, tf.text))
    elif ext in ('.txt', '.md'):
        # 🟥 한국어 원고는 CP949 로 저장된 것이 흔하다. UTF-8(BOM 허용) → CP949 순으로 시도하고, 둘 다
        #    실패하거나 파일을 못 열면 트레이스백(exit 1 = «발견」으로 오독)이 아니라 «못 읽음」이다
        #    (codex 교차 리뷰 B1).
        try:
            with open(path, 'rb') as fh:
                raw = fh.read()
        except OSError as e:
            raise ScriptUnreadable('원고 파일을 못 연다: %s' % type(e).__name__)
        text = None
        for enc in ('utf-8-sig', 'cp949'):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ScriptUnreadable('원고 인코딩을 못 읽는다(UTF-8 · CP949 둘 다 실패)')
        units = [(i, l) for i, l in enumerate(text.splitlines(), 1)]
    else:
        raise ScriptUnreadable('원고 형식을 모른다(%s) — .pptx(노트) · .txt · .md' % ext)
    out = []
    for u, text in units:
        for l in text.splitlines():
            l = l.strip()
            if not l or re.match(r'^\[(간지|진행)\]', l):
                continue
            out.append((u, re.sub(r'^\(이어서\)\s*', '', l)))
    if not out:
        raise ScriptUnreadable('낭독 줄 0 — 노트가 비었다')
    return out


def norm_chars(s):
    return [c for c in s if re.match(r'[가-힣A-Za-z0-9]', c)]


def nsyl(w):
    h = len(re.findall(r'[가-힣]', w))
    up = len(re.findall(r'[A-Z]', w))
    lo = len(re.findall(r'[a-z]+', w))
    dg = len(re.findall(r'[0-9]', w))
    return max(1, h + round(up * 1.5) + lo * 2 + round(dg * 1.5))


def jong(ch):
    return (ord(ch) - 0xAC00) % 28 if '가' <= ch <= '힣' else -1


# ── 항목별 계기 ───────────────────────────────────────────────────────────────────────────────
def row_pron(script, words, chunks):
    """① 원고 ↔ 받아쓰기 글자 대조. 1~2글자 치환 · 관형형 «-는」 탈락을 «후보」로 낸다."""
    sc, s_unit = [], []
    sw = []
    for u, l in script:
        for w in l.split():
            sw.append((u, w))
            for c in norm_chars(w):
                sc.append(c)
                s_unit.append(u)
    tc, t_chunk = [], []
    for w in words:
        for c in norm_chars(w['w']):
            tc.append(c)
            t_chunk.append(w['chunk'])
    if not sc or not tc:
        return None
    sm = difflib.SequenceMatcher(None, sc, tc, autojunk=False)
    ops = sm.get_opcodes()
    eq = set()
    s2t = {}
    for tag, i1, i2, j1, j2 in ops:
        if tag == 'equal':
            eq.update(range(i1, i2))
            for k in range(i2 - i1):
                s2t[i1 + k] = j1 + k
        else:
            for k in range(i1, i2):
                s2t[k] = min(j1, len(tc) - 1)

    def t_at(tj):
        return chunks[t_chunk[min(tj, len(t_chunk) - 1)]][0]

    subs = []
    for tag, i1, i2, j1, j2 in ops:
        a, b = ''.join(sc[i1:i2]), ''.join(tc[j1:j2])
        if tag == 'replace' and len(a) <= 2 and len(b) <= 2 and not re.search(r'[0-9A-Za-z]', a + b):
            subs.append({'t': t_at(j1), 'unit': s_unit[i1], 'script': a, 'heard': b,
                         'ctx': ''.join(sc[max(0, i1 - 4):i1]) + '[' + a + '→' + b + ']' + ''.join(sc[i2:i2 + 4])})
    neun = {'verb': [0, []], 'pron': [0, []]}
    pos = 0
    for u, w in sw:
        chars = norm_chars(w)
        core = re.sub(r'[^가-힣]', '', w)
        if len(core) >= 2 and core.endswith('는') and jong(core[-2]) == 0:
            kind = 'pron' if core[:-1] in PRONOUN_TOPIC else 'verb'
            ni = pos + max(i for i, c in enumerate(chars) if c == '는')   # 원고 글자열에서 그 «는」의 자리
            neun[kind][0] += 1
            if ni not in eq:
                neun[kind][1].append({'t': t_at(s2t.get(ni, 0)), 'unit': u, 'word': w})
        pos += len(chars)
    return {'ratio': round(sm.ratio(), 3), 'subs': subs, 'neun': neun}


def row_conj(words, chunks):
    """② 접속사: 덩어리 첫 낱말인가 · 홀로 섰나 · 앞뒤 쉼 · 늘임(1.0 = 지역 말속도 그대로)."""
    by = {}
    for i, w in enumerate(words):
        by.setdefault(w['chunk'], []).append(i)

    def gap_before(j):
        return round(chunks[j][0] - chunks[j - 1][1], 2) if j > 0 else None

    def gap_after(j):
        return round(chunks[j + 1][0] - chunks[j][1], 2) if j + 1 < len(chunks) else None

    def local_rate(j):
        js = [q for q in range(max(0, j - 3), min(len(chunks), j + 4)) if q in by]
        s = sum(nsyl(words[i]['w']) for q in js for i in by[q])
        d = sum(chunks[q][1] - chunks[q][0] for q in js)
        return s / d if d else None

    conj = []
    for i, w in enumerate(words):
        base = re.sub(r'[.,!?~"\']', '', w['w'])
        hit = next((c for c in CONJ if base == c), None)
        if not hit:
            continue
        j = w['chunk']
        ids = by[j]
        alone = len(ids) == 1
        first = ids[0] == i
        a, b = chunks[j]
        lr = local_rate(j)
        rec = {'conj': hit, 't': a if first else w['t0'], 'alone': alone, 'first_in_chunk': first,
               'pause_before': gap_before(j) if first else 0.0,
               'pause_after': gap_after(j) if (alone or ids[-1] == i) else 0.0}
        if alone:
            rec['dur'], rec['dur_src'] = round(b - a, 2), '음향 경계'
        else:
            rec['dur'], rec['dur_src'] = round(w['t1'] - w['t0'], 2), '토큰 시각(UNCALIBRATED)'
        # 1음절(«즉」)은 늘임을 안 낸다 — 음절 하나의 길이는 앞뒤 자음에 휘둘린다(프로토타입도 뺐다)
        rec['stretch'] = round((rec['dur'] / len(hit)) * lr, 2) if (lr and rec['dur'] > 0 and len(hit) >= 2) \
            else None
        conj.append(rec)
    commas, ends = [], []
    for i, w in enumerate(words):
        last_in_chunk = by[w['chunk']][-1] == i
        g = gap_after(w['chunk']) if last_in_chunk else 0.0
        if w['w'].endswith(','):
            commas.append(g)
        elif re.search(r'[.?!]["\']?$', w['w']) and g is not None:
            ends.append(g)
    return {'conj': conj, 'commas': commas, 'ends': ends}


def sentence_ends(chunks, words):
    """③ 측정 대상 끝: 받아쓰기가 있으면 문장부호로 끝나는 덩어리, 없으면 뒤 쉼이 긴 덩어리."""
    if words is not None:
        last = {}
        for w in words:
            last[w['chunk']] = w['w']
        return [j for j in sorted(last) if re.search(r'[.?!]["\']?$', last[j])], '받아쓰기 문장부호'
    out = []
    for j, (a, b) in enumerate(chunks):
        gap = chunks[j + 1][0] - b if j + 1 < len(chunks) else 99
        if gap >= C('phrase_end_gap_s') and b - a >= 0.5:
            out.append(j)
    return out, '뒤 쉼 ≥ %.2f초(UNCALIBRATED 대리)' % C('phrase_end_gap_s')


def chunk_level(db, a, b):
    seg = db[int(a * 100):int(b * 100)]
    return pct(seg, 90) if seg else None


def row_greet(db, chunks, words):
    """④ 첫·끝 덩어리 크기를 녹화 자기 분포(중앙 · 하위 10%)에 대고 잰다. 끝 덩어리는 끝 꺼짐까지."""
    lv = [chunk_level(db, a, b) for a, b in chunks]
    med, p10 = pct(lv, 50), pct(lv, 10)
    first, last = 0, len(chunks) - 1
    if words is not None:   # 받아쓰기가 있으면 인사 덩어리를 낱말로 찾는다(없으면 처음·끝 덩어리)
        text = {}
        for w in words:
            text[w['chunk']] = text.get(w['chunk'], '') + w['w']
        f = [j for j in sorted(text)[:5] if '안녕' in text[j]]
        l = [j for j in sorted(text)[-5:] if re.search('감사|고맙', text[j])]
        first, last = (f[0] if f else first), (l[-1] if l else last)
    out = {'median': med, 'p10': p10}
    for k, j in (('first', first), ('last', last)):
        a, b = chunks[j]
        e = {'chunk': j, 't': a, 'level': lv[j], 'vs_median': lv[j] - med, 'below_p10': lv[j] < p10,
             'rank_quiet': 1 + sum(1 for v in lv if v < lv[j]), 'n': len(lv)}
        if k == 'last':
            tw = min(C('tail_s'), (b - a) / 2)
            et = db[int((b - tw) * 100):int(b * 100)]
            eb = db[int(a * 100):int((b - tw) * 100)]
            if et and eb:
                e['tail_win_s'], e['tail_drop_db'] = tw, pct(et, 50) - pct(eb, 50)
        out[k] = e
    return out


# ── 코칭 본체 ─────────────────────────────────────────────────────────────────────────────────
def _row(n, name, status, cause=None, why='', obs=None, ev=None):
    return {'n': n, 'name': name, 'status': status, 'cause': cause, 'why': why,
            'obs': obs or [], 'evidence': ev or []}


NAMES = {1: '소리 내 읽기·발음', 2: '쉼표 여유 · 접속사 끌기', 3: '어미 밀당(문장 끝 억양)',
         4: '첫 인사·끝인사 강조', 5: '시선', 6: '스피치보다 내용'}


def coach(recording, script_path=None, skip=(), env=None):
    env = dict(os.environ if env is None else env)
    skip = set(skip)
    rep = {'lane': LANE, 'recording': recording, 'deps': {}, 'notes': [],
           'cal': {k: {'value': v, 'meaning': m, 'status': 'UNCALIBRATED'} for k, (v, m) in CAL.items()}}
    rows = {}
    for n in (5, 6):
        rows[n] = _row(n, NAMES[n], ST_NONE, CAUSE_DESIGN,
                       '오디오 계기 범위 밖 — 영상(시선 추정)이 필요하다' if n == 5 else
                       '계기 대상이 아니다 — 내용 정합은 이 스킬의 다른 레인(L1~L15)이 본다')
    for n in (1, 2, 3, 4):
        if n in skip:
            rows[n] = _row(n, NAMES[n], ST_NONE, CAUSE_DECLARED, '범위에서 뺌(선언 — --skip/--audio-only)')

    def finish(whole_reason=None):
        for n in (1, 2, 3, 4):
            if n not in rows:
                rows[n] = _row(n, NAMES[n], ST_NONE, CAUSE_DEP, whole_reason or '계기 미실행')
        if rep.get('decode_issue'):            # 들린 부분의 표는 내되 «다 들음」이라 하지 않는다
            for n in (1, 2, 3, 4):
                r_ = rows[n]
                if r_['cause'] in (None, CAUSE_INCOMPLETE):
                    r_['status'], r_['cause'] = ST_PARTIAL, CAUSE_INCOMPLETE
                    r_['why'] = '디코드 불완전: %s · %s' % (rep['decode_issue'], r_['why'])
        rep['rows'] = [rows[n] for n in range(1, 7)]
        unmeasured = [r['n'] for r in rep['rows'] if r['cause'] in (CAUSE_DEP, CAUSE_INCOMPLETE)]
        rep['rc'] = 2 if unmeasured else 0
        rep['unmeasured_rows'] = unmeasured
        return rep

    ffmpeg, why = find_tool(env, 'PREPREP_FFMPEG', 'ffmpeg')
    rep['deps']['ffmpeg'] = 'found' if ffmpeg else 'absent — ' + why
    if not ffmpeg:
        return finish('ffmpeg 없음 — 녹화를 못 연다(%s)' % why)
    try:                                       # 경계 E3 — 녹화 디코드(ffmpeg 출력 포함)
        if not os.path.isfile(recording):
            return finish('녹화 파일이 없다')
        x, dec_errs = decode(recording, ffmpeg, env)
    except Exception as e:   # noqa: BLE001 — 경계 E3
        return finish('디코드 실패(%s)' % _exc(e))
    # 부분 디코드 판정 — 오류 줄 · 길이 부족 · 길이 대조 불가, 셋 중 하나라도면 ①~④ 를 «부분」으로 강등
    issues = []
    if dec_errs:
        issues.append('디코드 오류 %d줄 — 첫 줄: %s' % (len(dec_errs), dec_errs[0][:100]))
    ffprobe, pwhy = find_tool(env, 'PREPREP_FFPROBE', 'ffprobe')
    if not ffprobe and not env.get('PREPREP_FFPROBE'):
        sib = os.path.join(os.path.dirname(ffmpeg), 'ffprobe')     # ffmpeg 옆에 같이 깔리는 게 보통이다
        if os.path.isfile(sib) and os.access(sib, os.X_OK):
            ffprobe, pwhy = sib, None
    decoded_s = len(x) / SR
    if not ffprobe:
        issues.append('녹화 길이 대조 못 함(%s) — 잘린 디코드를 못 가린다' % pwhy)
    else:
        try:                                   # 경계 E3b — ffprobe 출력
            dur = probe_duration(recording, ffprobe, env)
            rep['recording_s'] = round(dur, 2)
            if dur - decoded_s >= C('decode_short_frac') * dur:
                issues.append('디코드 길이 %.1f초 < 녹화 %.1f초 (%.1f%% 짧음)' % (decoded_s, dur,
                                                                       100 * (dur - decoded_s) / dur))
        except Exception as e:   # noqa: BLE001 — 경계 E3b
            issues.append('녹화 길이 대조 못 함(%s)' % _exc(e))
    rep['decode_issue'] = ' · '.join(issues) if issues else None
    db = frame_db(x)
    chunks, thr = segment(db)
    rep['duration_s'] = round(len(x) / SR, 2)
    rep['n_chunks'] = len(chunks)
    if not chunks:
        return finish('발화 덩어리 0 — 무음이거나 계기가 못 들었다(«발견 0」 아님)')
    rep['speech_s'] = round(sum(b - a for a, b in chunks), 2)

    # 받아쓰기(선택 의존성)
    words = None
    asr_gap = None
    need_asr = bool({1, 2} - skip)
    if need_asr:
        whisper, wwhy = find_tool(env, 'PREPREP_WHISPER_CLI', 'whisper-cli')
        model = env.get('PREPREP_WHISPER_MODEL')
        if not whisper:
            asr_why = 'whisper-cli 없음 (%s)' % wwhy
        elif not model:
            asr_why = 'PREPREP_WHISPER_MODEL 미지정 — 모델 경로 기본값은 일부러 없다'
        else:
            try:
                asr_why = check_model(model)
            except Exception as e:   # noqa: BLE001 — 경계 E2
                asr_why = '모델 경로 확인 실패(%s)' % _exc(e)
        rep['deps']['whisper'] = 'found' if not asr_why else 'absent — ' + asr_why
        if asr_why is None:
            tmp = None
            try:
                tmp = tempfile.mkdtemp(prefix='preprep_coach_')
                words, dropped, missing, quiet, miss_why = transcribe(x, chunks, whisper, model,
                                                                      env.get('PREPREP_WHISPER_LANG', 'ko'), tmp, env=env)
            except WhisperFailed as e:
                words, asr_why = None, str(e)
                rep['deps']['whisper'] = 'found but failed — ' + asr_why
            except OSError as e:     # 임시 폴더를 못 만든다
                words, asr_why = None, '받아쓰기 임시 폴더 실패(%s)' % _exc(e)
                rep['deps']['whisper'] = 'found but failed — ' + asr_why
            finally:
                if tmp:
                    shutil.rmtree(tmp, ignore_errors=True)
        if asr_why is None:
            rep['asr'] = {'chunks': len(chunks), 'missing': len(missing),
                          'missing_at': [fmt_t(chunks[j][0]) for j in missing],
                          'missing_why': {fmt_t(chunks[j][0]): miss_why.get(j, '') for j in missing},
                          'hallucination_dropped': len(dropped),
                          'hallucination_dropped_at': [fmt_t(chunks[j][0]) for j in dropped[:10]],
                          'unvoiced_empty': len(quiet)}
            if len(missing) == len(chunks):
                asr_why, words = '받아쓰기 출력 0 — whisper 가 한 덩어리도 못 냈다', None
            elif not words:
                # 🟥 파일은 나왔는데 낱말이 0 — 빈 받아쓰기를 «접속사 0개 · 발음 어긋남 0」으로 렌더하면
                #    «못 들었다」가 «깨끗하다」로 접힌다. 측정불가로 올린다.
                asr_why, words = '받아쓰기 낱말 0 — 빈 받아쓰기(또는 전부 환각으로 걸러짐)', None
            elif missing:
                # 🟥 노트만 달면 ② 가 «잡음」· exit 0 으로 나간다(codex 교차 리뷰 A1). 아래에서 ①② 를
                #    «부분 · 일부 못 들음」으로 묶고 rc 를 2 로 만든다 — 누락 비율과 위치를 같이 낸다.
                asr_gap = '받아쓰기 누락 덩어리 %d/%d (%.0f%%) — 위치 %s%s — 그 몫은 안 들었다' % (
                    len(missing), len(chunks), 100.0 * len(missing) / len(chunks),
                    ', '.join(rep['asr']['missing_at'][:10]), ' …' if len(missing) > 10 else '')
                rep['notes'].append(asr_gap)
        if asr_why:
            for n in (1, 2):
                if n not in rows:
                    rows[n] = _row(n, NAMES[n], ST_NONE, CAUSE_DEP, asr_why +
                                   ' — ③④ 는 오디오만으로 돈다. 범위에서 빼려면 --audio-only')

    # ① 발음
    if 1 not in rows:
        if not script_path:
            rows[1] = _row(1, NAMES[1], ST_NONE, CAUSE_DEP,
                           '원고 없음 — --script 로 원고(pptx 노트 · txt · md)를 주거나 --skip 1')
        else:
            r = None
            try:                                   # 경계 E4 — 원고 로드(pptx/txt/md)
                script = read_script(script_path)
            except ScriptUnreadable as e:
                rows[1] = _row(1, NAMES[1], ST_NONE, CAUSE_DEP, '원고를 못 읽는다: %s' % e)
            except Exception as e:   # noqa: BLE001 — 경계 E4
                rows[1] = _row(1, NAMES[1], ST_NONE, CAUSE_DEP, '원고를 못 읽는다(%s)' % _exc(e))
            if 1 not in rows:
                r = row_pron(script, words, chunks)   # 우리 분석 코드 — 경계 밖
            if 1 not in rows:
                if r is None:
                    rows[1] = _row(1, NAMES[1], ST_NONE, CAUSE_DEP, '대조할 글자가 0 — 원고나 받아쓰기가 비었다')
                else:
                    vlost, plost = r['neun']['verb'][1], r['neun']['pron'][1]
                    obs = ['원고 ↔ 받아쓰기 글자 일치율 %.3f' % r['ratio'],
                           '1~2글자 치환 %d곳(발음 «후보」)' % len(r['subs']),
                           '관형형 «-는」(열린 음절 + 는) 받아쓰기에서 사라짐: %d/%d · 대명사 주제 «-는」 %d/%d'
                           % (len(vlost), r['neun']['verb'][0], len(plost), r['neun']['pron'][0])]
                    ev = [(e['t'], '«%s」 → «-는」 탈락 후보 (원고 단위 %s)' % (e['word'], e['unit'])) for e in vlost[:3]]
                    ev += [(s['t'], '%s (원고 단위 %s)' % (s['ctx'], s['unit'])) for s in r['subs'][:3 - len(ev)]]
                    rows[1] = _row(1, NAMES[1], ST_PARTIAL, None,
                                   '후보다 — 받아쓰기 오류와 실제 발음을 이 계기는 못 가른다. 사람 귀로 확인 전엔 판정 아님',
                                   obs, ev)
                    rep['pron_candidates'] = {'subs': r['subs'], 'neun_lost': vlost, 'neun_pron_lost': plost}

    # ② 접속사 · 쉼표
    if 2 not in rows:
        r = row_conj(words, chunks)
        cj = r['conj']
        alone = [c for c in cj if c['alone'] and c['stretch'] is not None]
        drag = [c for c in alone if c['stretch'] >= C('drag_stretch')]
        obs = ['접속사 %d개 · 홀로 선 것 %d · 앞에 쉼 있는 것 %d'
               % (len(cj), sum(1 for c in cj if c['alone']), sum(1 for c in cj if (c['pause_before'] or 0) > 0))]
        if alone:
            s = sorted(c['stretch'] for c in alone)
            obs.append('홀로 선 접속사 늘임 %.2f~%.2f (중앙 %.2f) · 끌기 후보(≥%.2f, UNCALIBRATED) %d개'
                       % (s[0], s[-1], s[len(s) // 2], C('drag_stretch'), len(drag)))
        if r['commas']:
            rz = [g for g in r['commas'] if g and g > 0]
            obs.append('쉼표 %d곳 중 쉼으로 실현 %d · 쉼 중앙 %.2f초'
                       % (len(r['commas']), len(rz), pct(rz, 50) if rz else 0.0))
        if r['ends']:
            obs.append('문장 끝 뒤 쉼 중앙 %.2f초 (n=%d)' % (pct(r['ends'], 50), len(r['ends'])))
        pick = sorted(alone, key=lambda c: -c['stretch'])[:3] if alone else cj[:3]
        ev = [(c['t'], '«%s」 앞 쉼 %s초 · 뒤 쉼 %s초 · 늘임 %s (%s)'
               % (c['conj'], c['pause_before'], c['pause_after'], c['stretch'], c['dur_src'])) for c in pick]
        rows[2] = _row(2, NAMES[2], ST_MEASURED if cj else ST_PARTIAL, None,
                       '위치와 쉼은 음향 경계에서 잰다 · «끌기」 문턱은 UNCALIBRATED' if cj else
                       '접속사 0개 — 끌기는 못 봤다(0 은 «안 끌었다」가 아니다). 쉼표·문장 끝 쉼만',
                       obs, ev)
        rep['conjunctions'] = cj

    # 받아쓰기 일부 누락 → ①② 는 «부분」을 넘지 못하고 rc 를 0 으로 두지 않는다
    if asr_gap:
        for n in (1, 2):
            r_ = rows.get(n)
            if r_ is not None and r_['cause'] is None:
                r_['status'], r_['cause'] = ST_PARTIAL, CAUSE_INCOMPLETE
                r_['why'] = asr_gap + ' · ' + r_['why']

    # ③ 문장 끝
    if 3 not in rows:
        idx, src = sentence_ends(chunks, words)
        recs = []
        for j in idx:
            m = end_metrics(x, db, thr, chunks[j][0], chunks[j][1])
            if m.get('tail_st') is not None:
                recs.append(m)
        if len(recs) < 3:
            rows[3] = _row(3, NAMES[3], ST_NONE, CAUSE_DEP,
                           '피치를 잴 수 있는 문장 끝이 %d개(<3) — 유성 구간이 짧거나 피치를 못 잡았다' % len(recs))
        else:
            v = sorted(m['tail_st'] for m in recs)
            n = len(v)
            mean = sum(v) / n
            sd = math.sqrt(sum((z - mean) ** 2 for z in v) / n)
            pp = [m for m in recs if is_push_pull(m)]
            sus = [m for m in recs if octave_suspect(m)]
            obs = ['측정한 끝 %d개 (대상 고르기: %s)' % (n, src),
                   '끝 300ms 평균 피치(본문 대비) 중앙 %+.1f반음 · sd %.1f' % (v[n // 2], sd),
                   '하강(<-1) %d%% · 올림(>+1) %d%% · 올렸다 내림(밀당) %d%%'
                   % (round(100 * sum(1 for z in v if z < -1) / n), round(100 * sum(1 for z in v if z > 1) / n),
                      round(100 * len(pp) / n)),
                   '옥타브 오류 의심 %d개 — 분포에는 남기고 근거 예시·밀당 셈에서는 뺐다' % len(sus)]
            med = v[n // 2]
            rep_ = sorted([m for m in recs if not octave_suspect(m)], key=lambda m: abs(m['tail_st'] - med))
            ev = [(m['end'], '대표 끝(중앙 근처) %+.1f반음' % m['tail_st']) for m in rep_[:2]]
            ev += [(m['end'], '밀당 — 봉우리 %+.1f → 마지막 %+.1f' % (m['tail_peak_st'], m['tail_last_st']))
                   for m in pp[:1]]
            rows[3] = _row(3, NAMES[3], ST_PARTIAL, None,
                           '기준선 없음 — 한국어 평서문은 원래 내려간다. 모양의 묘사이지 결함 판정이 아니다', obs, ev)
            rep['sentence_ends'] = recs

    # ④ 첫·끝 인사
    if 4 not in rows:
        g = row_greet(db, chunks, words)
        obs = ['덩어리 크기(p90 dBFS) 중앙 %.1f · 하위 10%% %.1f%s' % (
            g['median'], g['p10'], ' — 덩어리 10개 미만이라 하위 10%는 곧 최솟값이다' if len(chunks) < 10 else '')]
        ev = []
        for tag, k in (('첫', 'first'), ('끝', 'last')):
            e = g[k]
            line = '%s 덩어리 %.1f dB (중앙 대비 %+.1f · 조용한 쪽부터 %d/%d)%s' % (
                tag, e['level'], e['vs_median'], e['rank_quiet'], e['n'],
                ' · 하위 10% 아래' if e['below_p10'] else '')
            if e.get('tail_drop_db') is not None:
                line += ' · 끝 %.2f초가 %+.1f dB' % (e['tail_win_s'], e['tail_drop_db'])
            obs.append(line)
            ev.append((e['t'], line))
        rows[4] = _row(4, NAMES[4], ST_MEASURED, None,
                       '녹화 자기 분포에 대고 잰 순위다 — 크다/작다의 절대 문턱은 없다', obs, ev)
    return finish()


# ── 출력 ─────────────────────────────────────────────────────────────────────────────────────
def render(rep):
    out = ['── 발표 코칭 (%s · advisory) ──' % rep['lane'],
           '   녹화: %s · 길이 %s · 발화 덩어리 %s'
           % (rep['recording'], fmt_t(rep.get('duration_s')), rep.get('n_chunks', '—')),
           '   의존성: ' + ' · '.join('%s %s' % kv for kv in rep['deps'].items())]
    if rep.get('decode_issue'):
        out.append('   ⚠️ 디코드 불완전 — %s' % rep['decode_issue'])
    if 'asr' in rep:
        a = rep['asr']
        out.append('   받아쓰기: 덩어리 %d · 누락 %d · 환각으로 버림 %d %s'
                   % (a['chunks'], a['missing'], a['hallucination_dropped'],
                      ('(' + ', '.join(a['hallucination_dropped_at']) + ')') if a['hallucination_dropped_at'] else ''))
        for t_, w_ in list(a.get('missing_why', {}).items())[:10]:
            out.append('     못 들음 %s — %s' % (t_, w_))
    out.append('   🟥 이 표는 판정이 아니다 — 숫자는 묘사이고 문턱은 전부 UNCALIBRATED 다.')
    out.append('')
    for r in rep['rows']:
        tag = r['status']
        if r['cause'] == CAUSE_DESIGN:
            tag += '(범위 밖 — 설계)'
        elif r['cause'] == CAUSE_DECLARED:
            tag += '(범위에서 뺌 — 선언)'
        elif r['cause'] == CAUSE_DEP:
            tag += '(UNMEASURED — 통과 아님)'
        elif r['cause'] == CAUSE_INCOMPLETE:
            tag += '(일부 못 들음 — 통과 아님)'
        out.append(' %s %s — %s' % ('①②③④⑤⑥'[r['n'] - 1], r['name'], tag))
        if r['why']:
            out.append('     ' + r['why'])
        for o in r['obs']:
            out.append('     · ' + o)
        for t, d in r['evidence']:
            out.append('     ▸ %s  %s' % (fmt_t(t), d))
    for n in rep['notes']:
        out.append('   ⚠️ ' + n)
    out.append('')
    out.append('   보정 상수 (전부 UNCALIBRATED — 녹화 1건 · TTS 대조군에서 고른 값):')
    for k, c in rep['cal'].items():
        out.append('     %-17s %-6s %s' % (k, c['value'], c['meaning']))
    if rep['rc'] == 2:
        out.append('\n   exit 2 — 측정 못 한 항목: %s. «발견 0」 이 아니라 «못 들었다」다(통과 아님).'
                   % ', '.join('①②③④⑤⑥'[n - 1] for n in rep['unmeasured_rows']))
    else:
        out.append('\n   exit 0 — 범위 안 항목을 전부 들었다. «좋은 발표」라는 뜻이 아니다.')
    return '\n'.join(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description='preprep L16 speech-coach — 리허설 녹화 코칭 6항 (advisory)')
    ap.add_argument('recording')
    ap.add_argument('--script', help='원고: 발표자 노트가 든 .pptx 또는 .txt/.md')
    ap.add_argument('--audio-only', action='store_true', help='①② 를 범위에서 뺀다(선언) — 받아쓰기 없이')
    ap.add_argument('--skip', default='', help='범위에서 뺄 항목 번호, 쉼표로(예: 1 또는 1,2)')
    ap.add_argument('--json', dest='json_out', help='전체 결과를 JSON 으로 쓴다')
    a = ap.parse_args(argv)
    # 🟥 isdigit() 는 «²» 같은 유니코드 숫자도 참이라 int() 에서 죽는다 — 허용 낱말을 정확히 대조한다.
    #    모르는 낱말은 argparse 오류(exit 2)로 — «측정 못 함」 과 같은 쪽, exit 1 크래시가 아니다.
    toks = [s.strip() for s in a.skip.split(',') if s.strip()]
    bad = [s for s in toks if s not in ('1', '2', '3', '4')]
    if bad:
        ap.error('--skip 은 1~4 만 받는다: %s' % ', '.join(repr(s) for s in bad))
    skip = {int(s) for s in toks}
    if a.audio_only:
        skip |= {1, 2}
    rep = coach(a.recording, a.script, skip)
    print(render(rep))
    if a.json_out:
        try:
            with open(a.json_out, 'w', encoding='utf-8') as fh:
                json.dump(rep, fh, ensure_ascii=False, indent=1)
        except OSError as e:
            print('JSON output could not be written: %s' % e, file=sys.stderr)
            return 2
    return rep['rc']


if __name__ == '__main__':
    try:
        from usage_ledger import observe as _observe
    except Exception:   # noqa: BLE001 — 원장은 옵트인이고 없어도 판정은 같다
        _observe = None
    sys.exit(_observe('preprep', 'speech-coach', main) if _observe else main())
